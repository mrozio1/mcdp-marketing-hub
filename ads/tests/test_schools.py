"""Private and charter schools from NCES near our offices: parsing, nearest office, charter filter, map data, export."""
import io

import pytest

from tests.test_scope import client  # noqa: F401  (pytest fixture)


def feat(attrs, x=None, y=None):
    return {"attributes": attrs, "geometry": {"x": x, "y": y} if x is not None else None}


def test_parse_and_charter_flag():
    from app import schools
    s = schools.parse(feat({"PPIN": "A123", "NAME": "ST. MARY CATHOLIC SCHOOL", "STREET": "100 MAIN ST", "CITY": "ROYAL OAK",
                            "STATE": "MI", "ZIP": "48067", "SCHOOLYEAR": "2023-2024"}, -83.14, 42.49), "school_private")
    assert s["key"] == "pss:A123" and s["name"] == "St. Mary Catholic School" and s["city"] == "Royal Oak" and s["lat"] == 42.49
    c = schools.parse(feat({"NCESSCH": "260000", "SCH_NAME": "Oak Academy", "LSTREET1": "1 Oak", "LCITY": "Troy", "LSTATE": "MI",
                            "LZIP": "48084", "GSLO": "KG", "GSHI": "08", "TOTAL": 412, "CHARTER_TEXT": "Yes"}, -83.1, 42.6), "school_charter")
    assert c["key"] == "ccd:260000" and c["grades"] == "K–8" and c["enrollment"] == 412
    assert schools.is_charter({"CHARTER_TEXT": "Yes"}) and schools.is_charter({"CHARTER_TEXT": "No"}) is False
    assert schools.is_charter({"NAME": "x"}) is None and schools.closed({"SY_STATUS_TEXT": "Closed"})
    assert schools.closed({"SY_STATUS_TEXT": "Open", "VIRTUAL": "Full Virtual"}) and not schools.closed({"SY_STATUS_TEXT": "Open", "VIRTUAL": "Supplemental Virtual"})
    assert schools.parse(feat({"NAME": "No id"}, 1, 2), "school_private") is None


def test_refresh_keeps_nearest_office_and_exports(client, monkeypatch):
    c, db, sync, mp = client
    from app import places, schools
    schools.init()
    places.add_office("Heights", "ChIJ_heights_office_01", 42.58, -83.03)
    places.add_office("Troy", "ChIJ_troy_office_0001", 42.60, -83.15)
    calls = []

    def near(url, lat, lng, miles, where="1=1", fields="*"):
        calls.append((url, where))
        if url == schools.PRIVATE_URL:
            return [feat({"PPIN": "P1", "NAME": "Shared Private", "CITY": "Troy", "STATE": "MI"}, -83.14, 42.60),
                    feat({"PPIN": "P2", "NAME": "Far Away Academy"}, -80.0, 40.0)]
        return [feat({"NCESSCH": "C1", "SCH_NAME": "Charter One", "CHARTER_TEXT": "Yes", "GSLO": "KG", "GSHI": "12"}, -83.04, 42.58),
                feat({"NCESSCH": "C2", "SCH_NAME": "Regular Public", "CHARTER_TEXT": "No"}, -83.04, 42.58),
                feat({"NCESSCH": "C3", "SCH_NAME": "Closed Charter", "CHARTER_TEXT": "Yes", "SY_STATUS_TEXT": "Closed"}, -83.04, 42.58)]
    monkeypatch.setattr(schools, "near", near)
    monkeypatch.setattr(schools.time, "sleep", lambda s: None)
    st = schools.refresh()
    assert st["private"] == 1 and st["charter"] == 1 and not st["error"]
    assert any(w == "CHARTER_TEXT='Yes'" for _, w in calls)
    troy = next(o["id"] for o in places.offices() if o["label"] == "Troy")
    rows = schools.listing()
    priv = next(r for r in rows if r["kind"] == "school_private")
    assert priv["office_id"] == troy and priv["miles"] < 1            # nearest office, not the first one asked
    monkeypatch.setattr(places, "browser_key", lambda: "test-key")
    data = c.get("/map/data.json").json()
    assert {s["name"] for s in data["schools"]} == {"Shared Private", "Charter One"}
    page = c.get("/map").text
    assert "Private schools" in page and "Charter schools" in page and "1 private and 1 charter" in page
    import openpyxl
    ws = openpyxl.load_workbook(io.BytesIO(c.get("/map/schools.xlsx").content)).active
    assert ws["D2"].value in ("Shared Private", "Charter One") and ws.max_row == 3


def test_missing_charter_field_is_reported(client, monkeypatch):
    c, db, sync, mp = client
    from app import places, schools
    schools.init()
    places.add_office("Solo", "ChIJ_solo_office_00001", 42.5, -83.0)

    def near(url, lat, lng, miles, where="1=1", fields="*"):
        if where != "1=1":
            raise schools.SchoolError("Invalid field")
        return [] if url == schools.PRIVATE_URL else [feat({"NCESSCH": "X", "SCH_NAME": "Mystery"}, -83.0, 42.5)]
    monkeypatch.setattr(schools, "near", near)
    monkeypatch.setattr(schools.time, "sleep", lambda s: None)
    with pytest.raises(schools.SchoolError):
        schools.refresh()
    assert "charter flag" in schools.status()["error"]
