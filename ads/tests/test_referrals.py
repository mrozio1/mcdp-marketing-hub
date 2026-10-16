"""Referral list: NPI Registry search, grouping into practices, geocoding, marks, export, routes."""
import csv
import importlib
import io
import json
import os
import threading
import urllib.parse
import zipfile
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

import pytest

PORT = 9967
OFFICE = {"id": "ChIJ_office_defiance01", "lat": 41.2845, "lng": -84.3558}
CALLS = []


def rec(npi, kind, name, a1, a2, z, codes, status="A", phone="419-555-0100"):
    basic = {"status": status}
    if kind == "org":
        basic["organization_name"] = name
    else:
        first, last, cred = name
        basic.update(first_name=first, last_name=last, credential=cred)
    return {"number": npi, "enumeration_type": "NPI-2" if kind == "org" else "NPI-1", "basic": basic,
            "addresses": [{"address_purpose": "MAILING", "address_1": "PO BOX 1", "city": "X", "state": "OH",
                           "postal_code": "43512"},
                          {"address_purpose": "LOCATION", "address_1": a1, "address_2": a2, "city": "DEFIANCE",
                           "state": "OH", "postal_code": z, "telephone_number": phone}],
            "taxonomies": [{"code": c, "desc": d, "primary": i == 0} for i, (c, d) in enumerate(codes)]}


GP = ("1223G0001X", "Dentist, General Practice")
ORTHO = ("1223X0400X", "Dentist, Orthodontics and Dentofacial Orthopedics")
PEDO = ("1223P0221X", "Dentist, Pediatric Dentistry")
HYG = ("124Q00000X", "Dental Hygienist")
NPI = {
    "43512": [
        rec("1", "org", "SMILE FAMILY DENTISTRY PC", "100 MAIN STREET", "SUITE 2", "435121234", [GP]),
        rec("2", "ind", ("JOHN", "DOE", "D.D.S."), "100 MAIN ST", "#2", "43512", [GP]),
        rec("3", "ind", ("ANN", "BRACE", "DMD"), "200 OAK AVE", "", "43512", [ORTHO]),            # ortho only: competition
        rec("4", "ind", ("BOB", "MIX", "DDS"), "300 ELM DR", "", "43512", [GP]),
        rec("5", "ind", ("CAT", "MIX", "DMD"), "300 ELM DRIVE", "", "43512", [ORTHO, GP]),        # mixed: in, flagged
        rec("6", "ind", ("HAL", "HYG", "RDH"), "400 PINE RD", "", "43512", [HYG]),                # hygienist: out
        rec("7", "org", "HEIGHTS ORTHODONTICS DENTAL", "1 OFFICE PL", "", "43512", [GP]),         # ours: out
        rec("8", "ind", ("OLD", "GONE", "DDS"), "500 ASH ST", "", "43512", [GP], status="D"),     # deactivated: out
        rec("9", "ind", ("KID", "SMILES", "DDS"), "600 BIRCH LN", "", "43512", [PEDO]),          # not geocoded -> ZIP
    ],
    "43545": [rec("10", "ind", ("FAR", "AWAY", "DDS"), "1 FAR RD", "", "43545", [GP])],          # outside the radius
}
GEO = {   # address -> lat, lng the stand-in Census geocoder returns
    "100 Main Street Suite 2": (41.2900, -84.3600), "200 Oak Ave": (41.2800, -84.3500),
    "300 Elm Dr": (41.3000, -84.3700), "1 Office Pl": (41.28451, -84.35581), "1 Far Rd": (41.5600, -84.3558),
}
ZCTA = "GEOID\tALAND\tAWATER\tALAND_SQMI\tAWATER_SQMI\tINTPTLAT\tINTPTLONG                                                                                                       \n" \
       "43512\t1\t1\t1\t1\t41.2845\t-84.3558\n43545\t1\t1\t1\t1\t41.3900\t-84.1300\n48187\t1\t1\t1\t1\t42.33\t-83.48\n"


class H(BaseHTTPRequestHandler):
    def log_message(self, *a):
        pass

    def _send(self, code, body, ctype="application/json"):
        data = body if isinstance(body, bytes) else (json.dumps(body) if not isinstance(body, str) else body).encode()
        self.send_response(code)
        self.send_header("Content-Type", ctype)
        self.end_headers()
        self.wfile.write(data)

    def do_GET(self):
        u = urllib.parse.urlparse(self.path)
        if u.path == "/gaz.zip":
            buf = io.BytesIO()
            with zipfile.ZipFile(buf, "w") as z:
                z.writestr("2024_Gaz_zcta_national.txt", ZCTA)
            return self._send(200, buf.getvalue(), "application/zip")
        q = dict(urllib.parse.parse_qsl(u.query))
        CALLS.append(q)
        if q.get("postal_code") == "99999":
            return self._send(200, {"Errors": [{"description": "Bad postal code"}]})
        assert q["version"] == "2.1" and q["address_purpose"] == "LOCATION" and q["taxonomy_description"] == "Dent*"
        res = NPI.get(q["postal_code"], [])
        skip, limit = int(q["skip"]), int(q["limit"])
        page = res[skip:skip + limit]
        return self._send(200, {"result_count": len(page), "results": page})

    def do_POST(self):
        body = self.rfile.read(int(self.headers["Content-Length"])).decode()
        assert "Public_AR_Current" in body and 'name="addressFile"' in body
        part = body.split("Content-Type: text/csv\r\n\r\n", 1)[1].rsplit("\r\n--", 1)[0]
        out = io.StringIO()
        w = csv.writer(out, quoting=csv.QUOTE_ALL)
        for key, street, city, state, z in csv.reader(io.StringIO(part)):
            if street in GEO:
                lat, lng = GEO[street]
                w.writerow([key, f"{street}, {city}, {state}, {z}", "Match", "Exact", "MATCHED ADDR",
                            f"{lng},{lat}", "1234", "L"])
            else:
                w.writerow([key, f"{street}, {city}, {state}, {z}", "No_Match"])
        return self._send(200, out.getvalue(), "text/csv")


@pytest.fixture(scope="module")
def env(tmp_path_factory):
    srv = ThreadingHTTPServer(("127.0.0.1", PORT), H)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    base = f"http://127.0.0.1:{PORT}"
    os.environ.update({"ADS_DB": str(tmp_path_factory.mktemp("db") / "ads.db"), "NPI_API_URL": base + "/api/",
                       "CENSUS_GEOCODER_URL": base + "/geo", "CENSUS_ZCTA_URL": base + "/gaz.zip",
                       "ADS_PASSWORD": "pw", "ADS_DISABLE_SCHEDULER": "1"})
    from app import connectors, db, places, referrals
    for m in (db, connectors, places, referrals):
        importlib.reload(m)
    referrals._sleep = lambda s: None
    referrals.NPI_PACE = 0
    referrals.NPI_PAGE = 2          # small pages, so paging is exercised
    db.init()
    places.init()
    referrals.init()
    db.set_setting("map_radius_miles", "15")
    db.set_setting("map_brand_terms", "Heights Orthodontics")
    places.add_office("Heights – Defiance", OFFICE["id"], OFFICE["lat"], OFFICE["lng"])
    with db.conn() as c:   # a Google pin right at Smile Family Dentistry
        c.execute("INSERT INTO map_places(place_id, layer, offers_ortho, office_id, miles, lat, lng, coords_at, ranked) "
                  "VALUES ('ChIJ_google_smile_0001','referrer',0,1,0.4,41.29002,-84.36001,?,1)", (db.now(),))
    yield db, places, referrals
    srv.shutdown()


def test_needs_offices(tmp_path, monkeypatch):
    from app import db, places, referrals
    monkeypatch.setattr(db, "DB_PATH", str(tmp_path / "empty.db"))
    db.init()
    places.init()
    referrals.init()
    with pytest.raises(referrals.RefError, match="offices"):
        referrals.refresh()


def test_build_list(env):
    db, places, referrals = env
    res = referrals.refresh()
    assert res["ok"] and res["zips"] == 2                        # 48187 is far away: not searched
    skips = [c["skip"] for c in CALLS if c.get("postal_code") == "43512"]
    assert skips == ["0", "2", "4", "6", "8"]                    # paged until a short page
    rows = {r["name"]: r for r in referrals.listing()}
    assert set(rows) == {"Smile Family Dentistry PC", "Dr. Bob Mix + 1 more", "Dr. Kid Smiles"}
    smile = rows["Smile Family Dentistry PC"]
    assert smile["dentists"] == ["John Doe, DDS"] and smile["address"] == "100 Main Street Suite 2"
    assert smile["place_id"] == "ChIJ_google_smile_0001" and smile["geo"] == "exact" and smile["phone"] == "419-555-0100"
    assert "query_place_id=ChIJ_google_smile_0001" in smile["maps_url"]
    mix = rows["Dr. Bob Mix + 1 more"]
    assert mix["has_ortho"] == 1 and mix["dentists"] == ["Bob Mix, DDS", "Cat Mix, DMD"]
    assert "Orthodontics" in mix["specialties"] and "General" in mix["specialties"]
    kid = rows["Dr. Kid Smiles"]
    assert kid["geo"] == "zip" and kid["specialties"] == "Pediatric Dentistry"     # placed at the ZIP centroid
    assert [r["name"] for r in referrals.listing()][0] == "Dr. Kid Smiles"           # nearest first (0 mi)
    assert all(r["office"] == "Heights – Defiance" for r in rows.values())
    comp = referrals.listing(kind="competitor")
    assert [r["name"] for r in comp] == ["Dr. Ann Brace"]                          # orthodontist-only practice
    assert comp[0]["specialties"] == "Orthodontics and Dentofacial Orthopedics" and comp[0]["geo"] == "exact"
    assert "Heights Orthodontics Dental" not in {r["name"] for r in comp}          # ours is in neither list


def test_marks_and_export(env):
    db, places, referrals = env
    rows = {r["name"]: r for r in referrals.listing()}
    assert referrals.mark(rows["Smile Family Dentistry PC"]["key"], starred=True, note="Dr Doe refers often")
    assert referrals.mark(rows["Dr. Kid Smiles"]["key"], omitted=True)
    assert not referrals.mark("nope", starred=True)
    import openpyxl
    data, n = referrals.export_xlsx()
    ws = openpyxl.load_workbook(io.BytesIO(data)).active
    body = list(ws.iter_rows(min_row=2, values_only=True))
    assert n == 2 and [r[3] for r in body] == ["Smile Family Dentistry PC", "Dr. Bob Mix + 1 more"]
    assert body[0][2] == "★" and body[0][11] == "Dr Doe refers often" and body[0][13] == "Yes"
    assert "query_place_id=ChIJ_google_smile_0001" in ws.cell(row=2, column=13).hyperlink.target
    assert "includes orthodontics" in body[1][5]
    data, n = referrals.export_xlsx(starred_only=True)
    assert n == 1
    data, n = referrals.export_xlsx(kind="competitor")
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert n == 1 and wb.active.title == "Competing practices" and wb.active.cell(row=2, column=4).value == "Dr. Ann Brace"
    assert "includes orthodontics" not in (wb.active.cell(row=2, column=6).value or "")
    # marks survive a rebuild
    referrals.refresh()
    again = {r["name"]: r for r in referrals.listing()}
    assert again["Smile Family Dentistry PC"]["starred"] == 1 and again["Dr. Kid Smiles"]["omitted"] == 1


def test_npi_error_is_reported(env, monkeypatch):
    db, places, referrals = env
    monkeypatch.setattr(referrals, "zips_near", lambda *a: ["99999"])
    with pytest.raises(referrals.RefError, match="Bad postal code"):
        referrals.refresh()
    st = referrals.status()
    assert st["ok"] is False and not st["busy"]


def test_routes(env):
    db, places, referrals = env
    referrals.refresh()
    from fastapi.testclient import TestClient
    from app import main
    importlib.reload(main)
    db.set_setting("maps_browser_key", "AIzaBrowser1")
    with TestClient(main.app) as cl:
        assert cl.get("/map/referrals.json").status_code == 401
        assert cl.post("/map/referrals/mark", json={"key": "x"}).status_code == 401
        assert cl.get("/map/referrals.xlsx", follow_redirects=False).status_code == 303
        cl.post("/login", data={"password": "pw"})
        html = cl.get("/map").text
        assert 'href="/map/list">Export Map List' in html and "ref-rows" not in html     # the list has its own page
        html = cl.get("/map/list").text
        office_id = places.offices()[0]["id"]
        assert 'id="ref-rows"' in html and "Update lists" in html and 'value="competitor"' in html
        assert f'<option value="{office_id}" selected>' in html                       # our practice first, not "All"
        assert '<option value="all" selected>' in cl.get("/map/list?office=all").text
        assert 'value="competitor" checked' in cl.get("/map/list?kind=competitor").text
        jc = cl.get(f"/map/referrals.json?office={office_id}&kind=competitor").json()
        assert jc["kind"] == "competitor" and [r["name"] for r in jc["rows"]] == ["Dr. Ann Brace"]
        j = cl.get("/map/referrals.json").json()
        assert j["radius"] == 15 and len(j["rows"]) == 3
        key = next(r["key"] for r in j["rows"] if r["name"] == "Dr. Bob Mix + 1 more")
        assert cl.post("/map/referrals/mark", json={"key": key, "starred": True, "note": "Call in May"}).json()["ok"]
        r = cl.get("/map/referrals.xlsx?office=&starred=1")
        assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxmlformats")
        assert "Referral-practices-all-offices" in r.headers["content-disposition"]
        r = cl.get(f"/map/referrals.xlsx?office={office_id}")
        assert "Referral-practices-Heights-Defiance" in r.headers["content-disposition"]
        r = cl.get(f"/map/referrals.xlsx?office={office_id}&kind=competitor")
        assert "Competing-practices-Heights-Defiance" in r.headers["content-disposition"]
        r = cl.post("/map/referrals/refresh", data={"back": "https://evil.example"}, follow_redirects=False)
        assert r.headers["location"] == "/map/list"
