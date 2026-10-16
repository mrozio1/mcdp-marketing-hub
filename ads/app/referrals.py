"""Map lists: dental practices (referral) and orthodontic practices (competition) near our offices,
from public government data.

Google's terms don't allow exporting names/addresses from the Places API, so the exportable list
comes from the NPI Registry (CMS's public registry of every licensed provider, free API, no
restriction on reuse), placed with the Census Bureau's free address geocoder and ZIP centroids.
Each practice is matched to its Google pin by location where one lines up (we keep only the
place ID, which Google allows). Omit / star / note marks are our own data and go into the export.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import os
import re
import threading
import time
import urllib.error
import urllib.parse
import urllib.request
import uuid
import zipfile
from collections import Counter
from datetime import datetime, timedelta

from . import connectors as cx
from . import db, places

NPI_URL = os.environ.get("NPI_API_URL", "https://npiregistry.cms.hhs.gov/api/")
GEOCODER_URL = os.environ.get("CENSUS_GEOCODER_URL",
                              "https://geocoding.geo.census.gov/geocoder/locations/addressbatch")
GAZETTEER_URL = os.environ.get("CENSUS_ZCTA_URL", "https://www2.census.gov/geo/docs/maps-data/data/gazetteer/"
                               "2024_Gazetteer/2024_Gaz_zcta_national.zip")
REFRESH_DAYS = 30
NPI_PACE = float(os.environ.get("NPI_PACE_SECONDS", "0.25"))
NPI_PAGE, NPI_MAX_SKIP = 200, 1000
GEOCODE_BATCH = 2000
SAME_PLACE_MILES = 0.08     # a Google pin this close to a practice's address is that practice
OUR_OFFICE_MILES = 0.05
ORTHO_CODE = "1223X0400X"
DENTAL_ORG_CODE = "261QD0000X"   # Clinic/Center, Dental

SCHEMA = """
CREATE TABLE IF NOT EXISTS zcta (zip TEXT PRIMARY KEY, lat REAL NOT NULL, lng REAL NOT NULL);
CREATE TABLE IF NOT EXISTS ref_practices (
  key TEXT PRIMARY KEY,
  name TEXT NOT NULL,
  dentists TEXT DEFAULT '[]',
  specialties TEXT DEFAULT '',
  address TEXT NOT NULL, city TEXT, state TEXT, zip TEXT,
  phone TEXT DEFAULT '',
  npis TEXT DEFAULT '',
  lat REAL, lng REAL, geo TEXT DEFAULT '',
  has_ortho INTEGER DEFAULT 0,
  ours INTEGER DEFAULT 0,
  place_id TEXT DEFAULT '',
  seen_at TEXT
);
CREATE TABLE IF NOT EXISTS ref_marks (
  key TEXT PRIMARY KEY,
  omitted INTEGER DEFAULT 0,
  starred INTEGER DEFAULT 0,
  note TEXT DEFAULT '',
  updated_at TEXT
);
"""

_lock = threading.Lock()
_started_at = datetime.now().isoformat(timespec="seconds")   # when this app process started
_sleep = time.sleep
_last = [0.0]


class RefError(RuntimeError):
    pass


KINDS = {"referrer": "Referral", "competitor": "Competition"}


def init():
    with db.conn() as c:
        c.executescript(SCHEMA)
        if "kind" not in {r["name"] for r in c.execute("PRAGMA table_info(ref_practices)")}:
            c.execute("ALTER TABLE ref_practices ADD COLUMN kind TEXT DEFAULT 'referrer'")


STALE_MINUTES = 20   # a running job updates its status at least every few minutes


def status():
    try:
        st = json.loads(db.get_setting("ref_status") or "{}")
    except ValueError:
        return {}
    if st.get("busy") and not _lock.locked():
        # Marked running, but nothing in this app is running it: the app restarted mid-way (an
        # update or reboot) or the work stopped without saying why. Clear it so it can be restarted.
        beat = st.get("beat") or st.get("started") or ""
        if not beat or beat < (datetime.now() - timedelta(minutes=STALE_MINUTES)).isoformat(timespec="seconds") \
                or _started_at > beat:
            st = {"busy": False, "ok": False, "at": db.now(),
                   "error": "It stopped part-way, most likely because the app restarted (an update or a reboot). "
                            "Everything found so far is kept; click Update lists to carry on."}
            db.set_setting("ref_status", json.dumps(st))
    return st


def _set_status(**kw):
    if kw.get("busy"):
        kw["beat"] = db.now()
    db.set_setting("ref_status", json.dumps(kw))


# ------------------------------------------------------------------ HTTP helpers
def _raw(url, data=None, headers=None, timeout=120):
    req = urllib.request.Request(url, data=data, headers=headers or {})
    try:
        with urllib.request.urlopen(req, timeout=timeout) as r:
            return r.read()
    except urllib.error.HTTPError as e:
        raise RefError(f"{urllib.parse.urlparse(url).netloc} answered HTTP {e.code}.")
    except urllib.error.URLError as e:
        raise RefError(f"Couldn't reach {urllib.parse.urlparse(url).netloc}: {e.reason}")


def _npi(params):
    wait = _last[0] + NPI_PACE - time.monotonic()
    if wait > 0:
        _sleep(wait)
    for attempt in range(4):
        _last[0] = time.monotonic()
        try:
            j = cx._http(NPI_URL + "?" + urllib.parse.urlencode(params), timeout=60)
            break
        except cx.ConnectorError as e:
            if attempt == 3:
                raise RefError(f"NPI Registry: {e}")
            _sleep(5 * (attempt + 1))
    if isinstance(j, dict) and j.get("Errors"):
        raise RefError("NPI Registry: " + "; ".join(str(x.get("description", x)) for x in j["Errors"]))
    return j.get("results") or []


# ------------------------------------------------------------------ ZIP centroids (Census gazetteer)
def ensure_zcta():
    with db.conn() as c:
        if c.execute("SELECT 1 FROM zcta LIMIT 1").fetchone():
            return
    data = _raw(GAZETTEER_URL, timeout=300)
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            name = next(n for n in z.namelist() if n.lower().endswith(".txt"))
            text = z.read(name).decode("latin-1")
    except (zipfile.BadZipFile, StopIteration):
        raise RefError("The Census ZIP code file couldn't be read.")
    rows = list(csv.reader(io.StringIO(text), delimiter="\t"))
    head = [h.strip().upper() for h in rows[0]]
    gi, la, lo = head.index("GEOID"), head.index("INTPTLAT"), head.index("INTPTLONG")
    out = []
    for r in rows[1:]:
        try:
            out.append((r[gi].strip(), float(r[la]), float(r[lo])))
        except (ValueError, IndexError):
            continue
    with db.conn() as c:
        c.executemany("INSERT OR REPLACE INTO zcta VALUES (?,?,?)", out)


def zips_near(lat, lng, miles):
    s, w, n, e = places.box_around(lat, lng, miles + 2)
    with db.conn() as c:
        rows = c.execute("SELECT * FROM zcta WHERE lat BETWEEN ? AND ? AND lng BETWEEN ? AND ?", (s, n, w, e))
        return [r["zip"] for r in rows if places.miles_between(lat, lng, r["lat"], r["lng"]) <= miles + 2]


def zip_point(z):
    with db.conn() as c:
        r = c.execute("SELECT lat, lng FROM zcta WHERE zip=?", (z[:5],)).fetchone()
        return (r["lat"], r["lng"]) if r else (None, None)


# ------------------------------------------------------------------ grouping NPI records into practices
ABBR = {"ROAD": "RD", "STREET": "ST", "AVENUE": "AVE", "DRIVE": "DR", "BOULEVARD": "BLVD", "LANE": "LN",
        "COURT": "CT", "PARKWAY": "PKWY", "HIGHWAY": "HWY", "PLACE": "PL", "CIRCLE": "CIR", "TERRACE": "TER",
        "NORTH": "N", "SOUTH": "S", "EAST": "E", "WEST": "W", "SUITE": "#", "STE": "#", "UNIT": "#",
        "BLDG": "#", "BUILDING": "#", "APT": "#", "ROOM": "#", "RM": "#"}


def norm_address(a1, a2=""):
    words = re.findall(r"[A-Z0-9#]+", f"{a1} {a2}".upper().replace("#", " # "))
    out = []
    for w in words:
        w = ABBR.get(w, w)
        if w == "#" and out and out[-1] == "#":
            continue
        out.append(w)
    return " ".join(out).replace("# ", "#")


def practice_key(a1, a2, zip5):
    return hashlib.sha1(f"{norm_address(a1, a2)}|{zip5}".encode()).hexdigest()[:16]


def _title(s):
    s = " ".join((s or "").split())
    if s.isupper():
        keep = {"DDS", "DMD", "MS", "MSD", "PC", "PLLC", "LLC", "PA", "II", "III", "IV", "USA", "MI", "OH", "IN"}
        s = " ".join(w if w.strip(".,") in keep else w.capitalize() for w in s.split())
    return s


def _location(rec):
    return next((a for a in rec.get("addresses") or [] if a.get("address_purpose") == "LOCATION"), None)


def _dental(rec):
    codes = [t.get("code", "") for t in rec.get("taxonomies") or []]
    return any(c.startswith("1223") or c == DENTAL_ORG_CODE for c in codes)


def group_records(records):
    """NPI records (dentists and dental organisations) -> practices keyed by normalised address."""
    groups = {}
    for rec in records:
        if not _dental(rec) or (rec.get("basic") or {}).get("status", "A") != "A":
            continue
        loc = _location(rec)
        if not loc or not loc.get("address_1"):
            continue
        z5 = (loc.get("postal_code") or "")[:5]
        key = practice_key(loc["address_1"], loc.get("address_2", ""), z5)
        g = groups.setdefault(key, {"recs": [], "loc": loc, "zip": z5})
        g["recs"].append(rec)
    out = {}
    for key, g in groups.items():
        orgs, people, specs, phones, npis = [], [], set(), Counter(), []
        primary_codes = []
        for rec in g["recs"]:
            b = rec.get("basic") or {}
            npis.append(str(rec.get("number", "")))
            tax = rec.get("taxonomies") or []
            prim = next((t for t in tax if t.get("primary")), tax[0] if tax else {})
            primary_codes.append(prim.get("code", ""))
            for t in tax:
                d = (t.get("desc") or "").replace("Dentist, ", "").strip()
                if t.get("code", "").startswith("1223") and d:
                    specs.add("General" if d in ("Dentist", "General Practice") else d)
            ph = (_location(rec) or {}).get("telephone_number")
            if ph:
                phones[ph] += 1
            if rec.get("enumeration_type") == "NPI-2" and b.get("organization_name"):
                orgs.append(_title(b["organization_name"]))
            elif b.get("last_name"):
                cred = (b.get("credential") or "").replace(".", "").upper()
                people.append(f"{_title(b.get('first_name', ''))} {_title(b['last_name'])}" + (f", {cred}" if cred else ""))
        people = sorted(set(people))
        has_general = any(c != ORTHO_CODE for c in primary_codes if c.startswith("1223") or c == DENTAL_ORG_CODE)
        has_ortho = ORTHO_CODE in primary_codes or any(
            t.get("code") == ORTHO_CODE for r in g["recs"] for t in r.get("taxonomies") or [])
        kind = "referrer" if has_general else "competitor"   # orthodontists only: a competing practice
        if orgs:
            name = sorted(set(orgs), key=len)[0]
        elif people:
            name = f"Dr. {people[0].split(',')[0]}" + (f" + {len(people) - 1} more" if len(people) > 1 else "")
        else:
            continue
        loc = g["loc"]
        addr = _title(" ".join(x for x in (loc.get("address_1"), loc.get("address_2")) if x))
        out[key] = {"key": key, "name": name, "dentists": json.dumps(people), "specialties": ", ".join(sorted(specs)),
                    "address": addr, "city": _title(loc.get("city", "")), "state": loc.get("state", ""),
                    "zip": g["zip"], "phone": phones.most_common(1)[0][0] if phones else "",
                    "npis": " ".join(sorted(set(npis))), "has_ortho": int(has_ortho), "kind": kind}
    return out


# ------------------------------------------------------------------ geocoding (Census batch)
def geocode(rows):
    """{key: (street, city, state, zip)} -> {key: (lat, lng)} for the addresses the Census geocoder matches."""
    found = {}
    items = list(rows.items())
    for i in range(0, len(items), GEOCODE_BATCH):
        buf = io.StringIO()
        w = csv.writer(buf)
        for key, (street, city, state, z) in items[i:i + GEOCODE_BATCH]:
            w.writerow([key, street, city, state, z])
        boundary = uuid.uuid4().hex
        body = (f"--{boundary}\r\nContent-Disposition: form-data; name=\"benchmark\"\r\n\r\nPublic_AR_Current\r\n"
                f"--{boundary}\r\nContent-Disposition: form-data; name=\"addressFile\"; filename=\"a.csv\"\r\n"
                f"Content-Type: text/csv\r\n\r\n{buf.getvalue()}\r\n--{boundary}--\r\n").encode()
        try:
            out = _raw(GEOCODER_URL, data=body, headers={"Content-Type": f"multipart/form-data; boundary={boundary}"},
                       timeout=600).decode("utf-8", errors="replace")
        except RefError:
            continue   # fall back to ZIP centroids for this batch
        for r in csv.reader(io.StringIO(out)):
            if len(r) < 3 or r[2].strip().lower() != "match":
                continue
            pt = next((x for x in r if re.fullmatch(r"\s*-?\d+\.\d+\s*,\s*-?\d+\.\d+\s*", x)), None)
            if pt:
                lng, lat = (float(v) for v in pt.split(","))
                found[r[0]] = (lat, lng)
    return found


# ------------------------------------------------------------------ refresh
def refresh(now=None):
    offs = [o for o in places.offices() if o["lat"] is not None]
    if not offs:
        raise RefError("Add our offices (and build the map once) first, so the list knows where to look.")
    if not _lock.acquire(blocking=False):
        raise RefError("The referral list is already being updated.")
    began = db.now()
    try:
        stamp = (now or datetime.now()).isoformat(timespec="seconds")
        _set_status(busy=True, started=began, step="Loading Census ZIP codes")
        ensure_zcta()
        miles = places.radius()
        zips = sorted({z for o in offs for z in zips_near(o["lat"], o["lng"], miles)})
        records = []
        for n, z in enumerate(zips, 1):
            if n % 10 == 0:
                _set_status(busy=True, started=began, step=f"Reading the NPI Registry: ZIP {n:,} of {len(zips):,}")
            for skip in range(0, NPI_MAX_SKIP + 1, NPI_PAGE):
                page = _npi({"version": "2.1", "postal_code": z, "address_purpose": "LOCATION",
                             "taxonomy_description": "Dent*", "limit": NPI_PAGE, "skip": skip})
                records += page
                if len(page) < NPI_PAGE:
                    break
        practices = group_records(records)
        with db.conn() as c:
            have = {r["key"]: dict(r) for r in c.execute("SELECT key, lat, lng, geo, address, zip FROM ref_practices")}
        to_geo = {k: (p["address"], p["city"], p["state"], p["zip"]) for k, p in practices.items()
                  if k not in have or have[k]["geo"] != "exact"}
        _set_status(busy=True, started=began, step=f"Placing {len(to_geo):,} addresses on the map (Census geocoder)")
        pts = geocode(to_geo) if to_geo else {}
        brands = places.brand_terms()
        pins = {kind: [p for p in places.places() if p["layer"] == kind and p["lat"] is not None] for kind in KINDS}
        rows = []
        for k, p in practices.items():
            if k in pts:
                lat, lng, geo = *pts[k], "exact"
            elif k in have and have[k]["geo"] == "exact":
                lat, lng, geo = have[k]["lat"], have[k]["lng"], "exact"
            else:
                lat, lng = zip_point(p["zip"])
                geo = "zip" if lat is not None else ""
            ours = int(any(b in p["name"].lower() for b in brands) or (geo == "exact" and any(
                places.miles_between(o["lat"], o["lng"], lat, lng) <= OUR_OFFICE_MILES for o in offs)))
            pid = ""
            refs = pins[p["kind"]]
            if geo == "exact" and refs:
                near = min(refs, key=lambda r: places.miles_between(r["lat"], r["lng"], lat, lng))
                if places.miles_between(near["lat"], near["lng"], lat, lng) <= SAME_PLACE_MILES:
                    pid = near["place_id"]
            rows.append((k, p["name"], p["dentists"], p["specialties"], p["address"], p["city"], p["state"], p["zip"],
                         p["phone"], p["npis"], lat, lng, geo, p["has_ortho"], ours, pid, stamp, p["kind"]))
        with db.conn() as c:
            c.executemany("INSERT OR REPLACE INTO ref_practices(key, name, dentists, specialties, address, city, state, "
                          "zip, phone, npis, lat, lng, geo, has_ortho, ours, place_id, seen_at, kind) "
                          "VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)", rows)
            c.execute("DELETE FROM ref_practices WHERE seen_at < ?", (stamp,))   # no longer registered there
        db.set_setting("ref_last_refresh", stamp)
        res = {"busy": False, "ok": True, "at": stamp, "practices": len(rows), "zips": len(zips),
               "referrers": sum(1 for r in rows if r[17] == "referrer" and not r[14]),
               "competitors": sum(1 for r in rows if r[17] == "competitor" and not r[14]),
               "exact": sum(1 for r in rows if r[12] == "exact"), "matched": sum(1 for r in rows if r[15])}
        _set_status(**res)
        return res
    except Exception as e:
        _set_status(busy=False, ok=False, at=db.now(), error=str(e)[:300])
        raise RefError(str(e))
    finally:
        _lock.release()


def refresh_in_background():
    def run():
        try:
            refresh()
        except RefError:
            pass
    threading.Thread(target=run, name="referrals", daemon=True).start()


def due(now=None):
    last = db.get_setting("ref_last_refresh")
    return bool(last) and (now or datetime.now()) - datetime.fromisoformat(last) >= timedelta(days=REFRESH_DAYS)


# ------------------------------------------------------------------ list, marks, export
def mark(key, omitted=None, starred=None, note=None):
    with db.conn() as c:
        if not c.execute("SELECT 1 FROM ref_practices WHERE key=?", (key,)).fetchone():
            return False
        c.execute("INSERT INTO ref_marks(key, updated_at) VALUES (?,?) ON CONFLICT(key) DO NOTHING", (key, db.now()))
        if omitted is not None:
            c.execute("UPDATE ref_marks SET omitted=?, updated_at=? WHERE key=?", (int(bool(omitted)), db.now(), key))
        if starred is not None:
            c.execute("UPDATE ref_marks SET starred=?, updated_at=? WHERE key=?", (int(bool(starred)), db.now(), key))
        if note is not None:
            c.execute("UPDATE ref_marks SET note=?, updated_at=? WHERE key=?", (note.strip()[:1000], db.now(), key))
    return True


def listing(office_id=None, include_omitted=True, kind="referrer"):
    """Practices within the radius of one office (or of any office), nearest first, with our marks."""
    offs = [o for o in places.offices() if o["lat"] is not None]
    if office_id:
        offs = [o for o in offs if o["id"] == office_id]
    radius = places.radius()
    with db.conn() as c:
        rows = [dict(r) for r in c.execute(
            "SELECT p.*, COALESCE(m.omitted,0) AS omitted, COALESCE(m.starred,0) AS starred, "
            "COALESCE(m.note,'') AS note FROM ref_practices p LEFT JOIN ref_marks m ON m.key=p.key "
            "WHERE p.lat IS NOT NULL AND p.ours=0 AND COALESCE(p.kind,'referrer')=?", (kind,))]
    out = []
    for r in rows:
        if r["omitted"] and not include_omitted:
            continue
        best = None
        for o in offs:
            d = places.miles_between(o["lat"], o["lng"], r["lat"], r["lng"])
            if d <= radius and (best is None or d < best[1]):
                best = (o, d)
        if best:
            r["office_id"], r["office"], r["miles"] = best[0]["id"], best[0]["label"], round(best[1], 1)
            r["dentists"] = json.loads(r["dentists"] or "[]")
            r["maps_url"] = maps_url(r)
            out.append(r)
    out.sort(key=lambda r: (r["office"].lower() if not office_id else "", r["miles"]))
    return out


def maps_url(r):
    q = {"api": "1", "query": f"{r['name']}, {r['address']}, {r['city']}, {r['state']} {r['zip']}"}
    if r.get("place_id"):
        q["query_place_id"] = r["place_id"]
    return "https://www.google.com/maps/search/?" + urllib.parse.urlencode(q)


def export_xlsx(office_id=None, starred_only=False, kind="referrer"):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    rows = [r for r in listing(office_id, include_omitted=False, kind=kind) if r["starred"] or not starred_only]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Referral practices" if kind == "referrer" else "Competing practices"
    head = ["Nearest office", "Miles", "Star", "Practice", "Dentists", "Specialties", "Address", "City", "State",
            "ZIP", "Phone", "Note", "Google Maps", "On our map", "NPI numbers", "Location"]
    ws.append(head)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F7471")
    for r in rows:
        ws.append([r["office"], r["miles"], "★" if r["starred"] else "", r["name"], "; ".join(r["dentists"]),
                   r["specialties"] + (" (includes orthodontics)" if r["has_ortho"] and kind == "referrer" else ""), r["address"], r["city"],
                   r["state"], r["zip"], r["phone"], r["note"], "Open in Google Maps", "Yes" if r["place_id"] else "",
                   r["npis"], "Address" if r["geo"] == "exact" else "ZIP code area (approximate)"])
        link = ws.cell(row=ws.max_row, column=13)
        link.hyperlink = r["maps_url"]
        link.font = Font(color="0F7471", underline="single")
    widths = [30, 7, 6, 36, 40, 26, 34, 18, 7, 8, 15, 40, 20, 11, 26, 24]
    for i, wdt in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = wdt
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column in (4, 5, 12))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    src = wb.create_sheet("About")
    for line in ("Source: NPI Registry (CMS, public data), placed with the U.S. Census Bureau geocoder.",
                 ("General dental practices (potential referrers)" if kind == "referrer" else
                  "Orthodontic practices (competition)") +
                 f" within {places.radius():g} miles of our offices; practices left out on the list page are excluded.",
                 "Names and addresses are as registered with the NPI Registry and may be out of date.",
                 f"Exported {datetime.now():%Y-%m-%d %H:%M}."):
        src.append([line])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), len(rows)


def start_scheduler():
    def loop():
        time.sleep(300)
        while True:
            try:
                if due():
                    refresh()
            except Exception:
                pass
            time.sleep(24 * 3600)
    threading.Thread(target=loop, name="referrals-monthly", daemon=True).start()
