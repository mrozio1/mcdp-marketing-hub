"""Reviews & comments: one inbox for Google reviews and Facebook / Instagram comments, plus the
Google review scoreboard (monthly counts against per-practice goals).

Everything waiting lands in "Needs reply" and stays there until someone replies, likes, hides it or
marks it done here, or it gets answered on the platform itself. Replies always post as the business
(the practice's Google profile, Facebook Page or Instagram account); the hub records who sent them.

Google reviews come from the Business Profile API (the practices' own data, so they may be kept).
Until Google approves that API, monthly review counts can be typed in by hand.
"""
from __future__ import annotations

import io
import json
import os
import re
import threading
import time
from datetime import date, datetime, timedelta, timezone

from . import connectors as cx
from . import db, scope

SOURCES = (("google", "Google reviews"), ("facebook", "Facebook"), ("instagram", "Instagram"))
SOURCE_LABEL = {"google": "Google review", "facebook": "Facebook", "instagram": "Instagram"}
COMMENT_DAYS = 60          # comments are read on posts from the last 60 days
INBOX_BACKFILL_DAYS = 30   # when the inbox first starts, anything older than this isn't added
DEFAULT_GOAL = 4
DONE_SHOWN = 100

SCHEMA = """
CREATE TABLE IF NOT EXISTS gbp_locations (
  name TEXT PRIMARY KEY,
  title TEXT NOT NULL DEFAULT '',
  address TEXT DEFAULT '',
  account TEXT DEFAULT '',
  practice_id INTEGER,
  rating REAL,
  total INTEGER,
  synced_at TEXT,
  full_at TEXT
);
CREATE TABLE IF NOT EXISTS gbp_reviews (
  id TEXT PRIMARY KEY,
  location TEXT NOT NULL,
  author TEXT DEFAULT '',
  stars INTEGER,
  text TEXT DEFAULT '',
  created TEXT NOT NULL,
  updated TEXT,
  reply TEXT DEFAULT '',
  reply_at TEXT
);
CREATE INDEX IF NOT EXISTS gbp_reviews_loc ON gbp_reviews(location, created);
CREATE TABLE IF NOT EXISTS inbox (
  id INTEGER PRIMARY KEY,
  key TEXT NOT NULL UNIQUE,
  source TEXT NOT NULL,
  account_id TEXT NOT NULL,
  external_id TEXT NOT NULL,
  author TEXT DEFAULT '',
  stars INTEGER,
  text TEXT DEFAULT '',
  created TEXT NOT NULL,
  post_id TEXT DEFAULT '',
  post_text TEXT DEFAULT '',
  post_link TEXT DEFAULT '',
  link TEXT DEFAULT '',
  status TEXT NOT NULL DEFAULT 'open',
  outcome TEXT DEFAULT '',
  handled_by TEXT DEFAULT '',
  handled_at TEXT,
  reply TEXT DEFAULT '',
  draft TEXT DEFAULT '',
  first_seen TEXT NOT NULL
);
CREATE INDEX IF NOT EXISTS inbox_status ON inbox(status, created);
CREATE TABLE IF NOT EXISTS review_goals (
  practice_id INTEGER NOT NULL,
  month TEXT NOT NULL,
  goal INTEGER NOT NULL,
  PRIMARY KEY (practice_id, month)
);
CREATE TABLE IF NOT EXISTS review_manual (
  practice_id INTEGER NOT NULL,
  month TEXT NOT NULL,
  count INTEGER NOT NULL,
  PRIMARY KEY (practice_id, month)
);
"""


class ReviewsError(ValueError):
    pass


# Reporting a Google review for removal. Google's API can't do it, so the hub opens Google's own
# Reviews Management Tool and keeps track: reported, then removed by Google or still up.
REPORT_TOOL = "https://support.google.com/business/workflow/9945796"
REPORT_FOLLOWUP_DAYS = 7
REPORT_REASONS = {   # Google's report reasons, in the order its form lists them, with when to use each
    "off_topic": ("Off topic", "Not about an experience at this business"),
    "spam": ("Spam", "Fake, posted by a bot or someone who was never here, or repeated"),
    "conflict": ("Conflict of interest", "A competitor, a former employee, or someone with a stake"),
    "profanity": ("Profanity", "Swearing or sexually explicit language"),
    "harassment": ("Bullying or harassment", "Personal attacks on a named person"),
    "hate": ("Discrimination or hate speech", "Harmful language about a group"),
    "personal": ("Personal information", "Shares someone's address, phone or other private details"),
    "not_helpful": ("Not helpful", "Doesn't help anyone decide about this business"),
}
REPORT_COLS = {"report_reason": "TEXT DEFAULT ''", "report_note": "TEXT DEFAULT ''", "reported_by": "TEXT DEFAULT ''",
               "reported_at": "TEXT", "removed_at": "TEXT"}


def init():
    with db.conn() as c:
        c.executescript(SCHEMA)
        have = {r["name"] for r in c.execute("PRAGMA table_info(inbox)")}
        for col, ddl in REPORT_COLS.items():
            if col not in have:
                c.execute(f"ALTER TABLE inbox ADD COLUMN {col} {ddl}")


# ------------------------------------------------------------------ small helpers
def utcnow():
    return datetime.now(timezone.utc).replace(tzinfo=None, microsecond=0)


def _iso(dt):
    return dt.isoformat(timespec="seconds")


def _utc(s):
    """Any ISO timestamp (with or without zone) -> naive UTC 'YYYY-MM-DDTHH:MM:SS'."""
    if not s:
        return ""
    s = s.replace("Z", "+00:00").replace("+0000", "+00:00")
    try:
        d = datetime.fromisoformat(s)
    except ValueError:
        return s[:19]
    if d.tzinfo:
        d = d.astimezone(timezone.utc).replace(tzinfo=None)
    return _iso(d.replace(microsecond=0))


def _dt(s):
    try:
        return datetime.fromisoformat(s[:19])
    except (TypeError, ValueError):
        return None


def ago(s, now=None):
    d = _dt(s)
    if not d:
        return ""
    secs = max(0, ((now or utcnow()) - d).total_seconds())
    if secs < 3600:
        return f"{max(1, int(secs // 60))}m ago"
    if secs < 86400:
        return f"{int(secs // 3600)}h ago"
    days = int(secs // 86400)
    return "yesterday" if days == 1 else f"{days}d ago"


def hours_since(s, now=None):
    d = _dt(s)
    return ((now or utcnow()) - d).total_seconds() / 3600 if d else 0


def month_key(d):
    return f"{d.year:04d}-{d.month:02d}"


def months_back(n=12, today=None):
    """The last n month keys, oldest first, ending with the current month."""
    today = today or date.today()
    y, m, out = today.year, today.month, []
    for _ in range(n):
        out.append(f"{y:04d}-{m:02d}")
        m -= 1
        if m == 0:
            y, m = y - 1, 12
    return out[::-1]


def month_label(key, short=False):
    d = date(int(key[:4]), int(key[5:7]), 1)
    return d.strftime("%b" if short else "%b %Y")


def int_setting(key, default):
    try:
        return int(db.get_setting(key) or default)
    except ValueError:
        return default


def overdue_hours():
    return int_setting("inbox_overdue_hours", 48)


def check_minutes():
    return max(10, int_setting("inbox_check_minutes", 30))


# ------------------------------------------------------------------ practices and accounts
def practices():
    """The practices reviews are tracked for: every practice except the master ad record."""
    allp = db.practices()
    filtered = [p for p in allp if scope.has_filter(p)]
    if not filtered:
        return allp
    with_loc = {l["practice_id"] for l in locations()}
    return [p for p in allp if scope.has_filter(p) or db.social_ids(p, "fb") or db.social_ids(p, "ig")
            or p["id"] in with_loc]


def locations():
    with db.conn() as c:
        return [dict(r) for r in c.execute("SELECT * FROM gbp_locations ORDER BY title COLLATE NOCASE")]


def account_practice(plist=None):
    """{(source, account id): practice id} for Google locations, Facebook Pages and Instagram accounts."""
    out = {}
    for loc in locations():
        if loc["practice_id"]:
            out[("google", loc["name"])] = loc["practice_id"]
    for p in plist if plist is not None else db.practices():
        for a in db.social_ids(p, "fb"):
            out.setdefault(("facebook", a), p["id"])
        for a in db.social_ids(p, "ig"):
            out.setdefault(("instagram", a), p["id"])
    return out


def account_names():
    names = {}
    for key in ("fb_pages", "ig_accounts"):
        try:
            for it in json.loads(db.get_setting(key) or "[]"):
                names[it["id"]] = it.get("name") or it["id"]
        except (ValueError, KeyError, TypeError):
            continue
    for loc in locations():
        names[loc["name"]] = loc["title"]
    return names


def gbp_connected():
    return bool(db.get_setting("gbp_refresh_token")) and all(cx.gbp_client())


# ------------------------------------------------------------------ flags
SPAM_RX = re.compile(r"\b(followers?|dm (me|us)|promo(tion)?s?\b|crypto|bitcoin|forex|whatsapp|telegram|"
                     r"check (my|our) (bio|page)|free (likes|followers)|earn \$|cash ?app|investment|sugar ?(daddy|mommy))",
                     re.I)
UNHAPPY_RX = re.compile(r"\b(rude|terrible|horrible|awful|worst|disappoint\w*|unprofessional|refund|complain\w*|"
                        r"never (again|called|answered)|no ?one (answered|called)|waited|waiting (for|over)|"
                        r"ridiculous|scam|lied|charged|overcharg\w*|billing|frustrat\w*|angry|upset|unacceptable|"
                        r"called (\w+ )?(times|twice))\b", re.I)


def flags(item):
    """('spam' | 'unhappy' | 'overdue' | None) for an open item."""
    if item["status"] != "open":
        return None
    text = item.get("text") or ""
    if item["source"] != "google" and SPAM_RX.search(text):
        return "spam"
    if (item["source"] == "google" and item.get("stars") and item["stars"] <= 2) or \
            (item["source"] != "google" and UNHAPPY_RX.search(text)):
        return "unhappy"
    if hours_since(item["created"]) > overdue_hours():
        return "overdue"
    return None


FLAG_TEXT = {"spam": "Likely spam", "unhappy": "Unhappy", "overdue": "Overdue"}


def no_reply_hint(item):
    """A comment that only tags friends (\"@sam look\") usually needs a like, not a reply."""
    t = (item.get("text") or "").strip()
    return item["source"] != "google" and t.startswith("@") and "?" not in t and len(t.split()) <= 8


# ------------------------------------------------------------------ inbox storage
def _inbox_start():
    v = db.get_setting("inbox_start")
    if not v:
        v = _iso(utcnow() - timedelta(days=INBOX_BACKFILL_DAYS))
        db.set_setting("inbox_start", v)
    return v


def upsert_items(source, account_id, items, outside_label):
    """Add new reviews / comments to the inbox; close open ones that were answered on the platform.
    items: [{'id','author','stars','text','created','post_id','post_text','post_link','link','replied','hidden'}]
    Returns how many were new."""
    start = _inbox_start()
    now = _iso(utcnow())
    new = 0
    with db.conn() as c:
        for it in items:
            key = f"{source}:{it['id']}"
            created = _utc(it["created"])
            row = c.execute("SELECT id, status FROM inbox WHERE key=?", (key,)).fetchone()
            answered = bool(it.get("replied")) or bool(it.get("hidden"))
            if row is None:
                if created < start:
                    continue
                new += 1
                outcome = "outside" if it.get("replied") else ("hidden" if it.get("hidden") else "")
                c.execute("INSERT INTO inbox(key, source, account_id, external_id, author, stars, text, created, post_id, "
                          "post_text, post_link, link, status, outcome, handled_by, handled_at, reply, first_seen) "
                          "VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                          (key, source, account_id, it["id"], it.get("author") or "", it.get("stars"), it.get("text") or "",
                           created, it.get("post_id") or "", it.get("post_text") or "", it.get("post_link") or "",
                           it.get("link") or "", "done" if answered else "open", outcome,
                           outside_label if answered else "", now if answered else None,
                           it.get("reply_text") or "", now))
            else:
                c.execute("UPDATE inbox SET author=?, stars=?, text=?, post_text=?, post_link=?, link=? WHERE id=?",
                          (it.get("author") or "", it.get("stars"), it.get("text") or "", it.get("post_text") or "",
                           it.get("post_link") or "", it.get("link") or "", row["id"]))
                if row["status"] == "open" and answered:
                    c.execute("UPDATE inbox SET status='done', outcome=?, handled_by=?, handled_at=?, reply=? WHERE id=?",
                              ("outside" if it.get("replied") else "hidden", outside_label, now,
                               it.get("reply_text") or "", row["id"]))
    return new


def item(iid):
    with db.conn() as c:
        r = c.execute("SELECT * FROM inbox WHERE id=?", (iid,)).fetchone()
        return dict(r) if r else None


def _decorate(rows, owners, plist, names, now):
    by_id = {p["id"]: p for p in plist}
    out = []
    for r in rows:
        pid = owners.get((r["source"], r["account_id"]))
        r["practice_id"] = pid
        r["practice_name"] = by_id[pid]["name"] if pid in by_id else "Not linked to a practice"
        r["account_name"] = names.get(r["account_id"], r["account_id"])
        r["source_label"] = SOURCE_LABEL[r["source"]]
        r["ago"] = ago(r["created"], now)
        r["handled_ago"] = ago(r["handled_at"], now) if r.get("handled_at") else ""
        r["flag"] = flags(r)
        r["hours"] = hours_since(r["created"], now)
        if r.get("reported_at"):
            r["report_label"] = REPORT_REASONS.get(r.get("report_reason") or "", ("Reported", ""))[0]
            r["report_days"] = int(hours_since(r["reported_at"], now) // 24)
            r["report_ago"] = ago(r["reported_at"], now)
            r["report_stale"] = r["status"] == "reported" and r["report_days"] >= REPORT_FOLLOWUP_DAYS
        out.append(r)
    return out


def inbox(practice_id=None, source="all", status="open", plist=None):
    """Items for the inbox list, already sorted: open = unhappy first, then oldest first;
    done = most recently handled first."""
    plist = plist if plist is not None else practices()
    owners = account_practice(db.practices())
    names = account_names()
    now = utcnow()
    with db.conn() as c:
        if status == "done":
            rows = [dict(r) for r in c.execute("SELECT * FROM inbox WHERE status='done' ORDER BY handled_at DESC, id DESC")]
        elif status == "reported":
            rows = [dict(r) for r in c.execute("SELECT * FROM inbox WHERE status='reported' ORDER BY reported_at, id")]
        else:
            rows = [dict(r) for r in c.execute("SELECT * FROM inbox WHERE status='open' ORDER BY created, id")]
    rows = _decorate(rows, owners, plist, names, now)
    if practice_id:
        rows = [r for r in rows if r["practice_id"] == practice_id]
    counts = {"all": len(rows), **{k: sum(1 for r in rows if r["source"] == k) for k, _ in SOURCES}}
    if source != "all":
        rows = [r for r in rows if r["source"] == source]
    if status == "done":
        rows = rows[:DONE_SHOWN]
    else:
        rows.sort(key=lambda r: (0 if r["flag"] == "unhappy" else 1, r["created"]))
    return rows, counts


def reported_count(practice_id=None, plist=None):
    try:
        return len(inbox(practice_id, "all", "reported", plist)[0])
    except Exception:
        return 0


def report(iid, reason, note, user_name):
    """Record that a Google review was reported for removal (the report itself is made in Google's tool)."""
    it = item(iid)
    if not it:
        raise ReviewsError("That review isn't in the inbox any more.")
    if it["source"] != "google":
        raise ReviewsError("Only Google reviews are reported this way. Hide Facebook and Instagram comments instead.")
    if reason not in REPORT_REASONS:
        raise ReviewsError("Pick the reason you gave Google.")
    if it["status"] == "reported":
        raise ReviewsError("This review is already marked as reported.")
    if it.get("removed_at"):
        raise ReviewsError("Google has already removed this review.")
    with db.conn() as c:
        c.execute("UPDATE inbox SET status='reported', report_reason=?, report_note=?, reported_by=?, reported_at=?, "
                  "removed_at=NULL WHERE id=?", (reason, (note or "").strip()[:1000], user_name, _iso(utcnow()), iid))
    return item(iid)


def unreport(iid):
    """Undo 'reported' (e.g. the report wasn't sent after all): back to Needs reply, or Done if it was answered."""
    it = item(iid)
    if not it or it["status"] != "reported":
        raise ReviewsError("This review isn't marked as reported.")
    with db.conn() as c:
        c.execute("UPDATE inbox SET status=?, report_reason='', report_note='', reported_by='', reported_at=NULL "
                  "WHERE id=?", ("done" if it["handled_at"] else "open", iid))
    return item(iid)


def _mark_removed(location, present_ids, complete):
    """After pulling a location's reviews: reported reviews Google no longer returns were removed. Only when the
    pull was complete (every page), so a short page never looks like a removal. Returns how many were removed."""
    if not complete:
        return 0
    now = _iso(utcnow())
    with db.conn() as c:
        gone = [r for r in c.execute("SELECT id, external_id FROM inbox WHERE source='google' AND account_id=? "
                                     "AND status='reported'", (location,)) if r["external_id"] not in present_ids]
        for r in gone:
            c.execute("UPDATE inbox SET status='done', outcome='removed', removed_at=?, handled_by='Removed by Google', "
                      "handled_at=? WHERE id=?", (now, now, r["id"]))
        stored = [r["id"] for r in c.execute("SELECT id FROM gbp_reviews WHERE location=?", (location,))]
        stale = [i for i in stored if i not in present_ids]
        if stale:   # removed reviews stop counting on the scoreboard
            c.executemany("DELETE FROM gbp_reviews WHERE id=?", [(i,) for i in stale])
    return len(gone)


def open_count(practice_id=None):
    try:
        if practice_id is None:
            with db.conn() as c:
                return c.execute("SELECT COUNT(*) FROM inbox WHERE status='open'").fetchone()[0]
        return len(inbox(practice_id)[0])
    except Exception:
        return 0


def inbox_stats(practice_id=None, plist=None):
    plist = plist if plist is not None else practices()
    owners = account_practice(db.practices())
    now = utcnow()
    with db.conn() as c:
        rows = [dict(r) for r in c.execute(
            "SELECT * FROM inbox WHERE status='open' OR handled_at >= ?", (_iso(now - timedelta(days=30)),))]
    if practice_id:
        rows = [r for r in rows if owners.get((r["source"], r["account_id"])) == practice_id]
    opens = [r for r in rows if r["status"] == "open"]
    week = _iso(now - timedelta(days=7))
    handled = [r for r in rows if r["status"] == "done" and (r["handled_at"] or "") >= week and r["outcome"] != "outside"]
    replied = sorted(hours_since(r["created"], _dt(r["handled_at"])) for r in rows
                     if r["status"] == "done" and r["outcome"] == "replied" and r["handled_at"])
    median = replied[len(replied) // 2] if replied else None
    oldest = max((hours_since(r["created"], now) for r in opens), default=None)
    return {"open": len(opens), "by_source": {k: sum(1 for r in opens if r["source"] == k) for k, _ in SOURCES},
            "oldest": oldest, "overdue": sum(1 for r in opens if hours_since(r["created"], now) > overdue_hours()),
            "handled": len(handled), "replies": sum(1 for r in handled if r["outcome"] == "replied"),
            "likes": sum(1 for r in handled if r["outcome"] == "liked"),
            "hidden": sum(1 for r in handled if r["outcome"] == "hidden"),
            "marked": sum(1 for r in handled if r["outcome"] == "done"),
            "median_hours": median}


def duration(hours):
    if hours is None:
        return "—"
    if hours < 1:
        return f"{max(1, round(hours * 60))}m"
    if hours < 48:
        return f"{round(hours)}h"
    return f"{round(hours / 24)}d"


# ------------------------------------------------------------------ tokens (kept in memory only)
_tok_lock = threading.Lock()
_TOKENS = {"at": 0.0, "facebook": {}, "instagram": {}, "ig_names": {}}


def _meta_tokens(force=False):
    with _tok_lock:
        if force or time.time() - _TOKENS["at"] > 900:
            pages, igs = cx.meta_pages()
            _TOKENS.update(at=time.time(), facebook={p["id"]: p.get("tokens") or ([p["token"]] if p["token"] else []) for p in pages},
                           instagram={i["id"]: i.get("tokens") or ([i["token"]] if i["token"] else []) for i in igs},
                           ig_names={i["id"]: i["name"] for i in igs})
        return _TOKENS


# ------------------------------------------------------------------ sync
_sync_lock = threading.Lock()


def sync_google(today=None):
    """Pull reviews for every Google location linked to a practice. All pages once a day (the
    scoreboard's history), otherwise just the newest page."""
    if not gbp_connected():
        return {}
    linked = [l for l in locations() if l["practice_id"]]
    if not linked:
        return {}
    try:
        client = cx.GBPClient(db.get_setting("gbp_refresh_token"))
    except cx.ConnectorError as e:
        return {l["name"]: {"ok": False, "error": str(e)} for l in linked}
    out, now = {}, utcnow()
    with db.conn() as c:
        watching = {r["account_id"] for r in c.execute("SELECT DISTINCT account_id FROM inbox WHERE source='google' "
                                                        "AND status='reported'")}
    for loc in linked:
        # a reported review is looked for on every check, so a removal shows within the half hour
        full = not loc["full_at"] or hours_since(loc["full_at"], now) > 20 or loc["name"] in watching
        try:
            revs, avg, total = client.reviews(loc["name"], all_pages=full)
        except cx.ConnectorError as e:
            out[loc["name"]] = {"ok": False, "error": str(e)}
            continue
        with db.conn() as c:
            c.executemany("INSERT INTO gbp_reviews(id, location, author, stars, text, created, updated, reply, reply_at) "
                          "VALUES (?,?,?,?,?,?,?,?,?) ON CONFLICT(id) DO UPDATE SET author=excluded.author, "
                          "stars=excluded.stars, text=excluded.text, updated=excluded.updated, reply=excluded.reply, "
                          "reply_at=excluded.reply_at",
                          [(r["id"], loc["name"], r["author"], r["stars"], r["text"], _utc(r["created"]),
                            _utc(r["updated"]), r["reply"], _utc(r["reply_at"])) for r in revs])
            c.execute("UPDATE gbp_locations SET rating=COALESCE(?, rating), total=COALESCE(?, total), synced_at=?"
                      + (", full_at=?" if full else "") + " WHERE name=?",
                      (avg, total, _iso(now), *((_iso(now),) if full else ()), loc["name"]))
        new = upsert_items("google", loc["name"], [{**r, "replied": bool(r["reply"]), "reply_text": r["reply"]}
                                                   for r in revs], "Answered on Google")
        complete = full and total is not None and len({r["id"] for r in revs}) >= int(total)
        removed = _mark_removed(loc["name"], {r["id"] for r in revs}, complete)
        out[loc["name"]] = {"ok": True, "reviews": len(revs), "new": new, "removed": removed}
    return out


def sync_comments(today=None):
    plist = db.practices()
    fb = sorted({a for p in plist for a in db.social_ids(p, "fb")})
    ig = sorted({a for p in plist for a in db.social_ids(p, "ig")})
    if not (fb or ig) or not cx.meta_configured():
        return {}
    try:
        toks = _meta_tokens(force=True)
    except cx.ConnectorError as e:
        return {a: {"ok": False, "error": str(e)} for a in fb + ig}
    since = (today or date.today()) - timedelta(days=COMMENT_DAYS)
    out = {}
    for pid in fb:
        tokens = toks["facebook"].get(pid) or []
        if not tokens:
            out[pid] = {"ok": False, "error": "No access to this Facebook Page's comments yet."}
            continue
        try:
            comments, _ = cx.facebook_comments(pid, tokens, since)
            new = upsert_items("facebook", pid, [{**c, "replied": c["page_replied"]} for c in comments], "Answered on Facebook")
            out[pid] = {"ok": True, "comments": len(comments), "new": new}
        except cx.ConnectorError as e:
            out[pid] = {"ok": False, "error": str(e)}
    for iid in ig:
        tokens = toks["instagram"].get(iid) or []
        if not tokens:
            out[iid] = {"ok": False, "error": "No access to this Instagram account's comments yet."}
            continue
        try:
            comments, _ = cx.instagram_comments(iid, toks["ig_names"].get(iid, ""), tokens, since)
            new = upsert_items("instagram", iid, [{**c, "replied": c["page_replied"]} for c in comments],
                               "Answered on Instagram")
            out[iid] = {"ok": True, "comments": len(comments), "new": new}
        except cx.ConnectorError as e:
            out[iid] = {"ok": False, "error": str(e)}
    return out


def sync_inbox(today=None):
    if not _sync_lock.acquire(blocking=False):
        return {"busy": True}
    try:
        res = {"google": sync_google(today), "meta": sync_comments(today)}
        names = account_names()
        errors = [f"{names.get(k, k)}: {v['error']}" for part in res.values() for k, v in part.items() if not v.get("ok")]
        status = {"at": db.now(), "new": sum(v.get("new", 0) for part in res.values() for v in part.values()),
                  "accounts": sum(len(p) for p in res.values()), "errors": errors[:8], "error_count": len(errors)}
        db.set_setting("inbox_status", json.dumps(status))
        return status
    finally:
        _sync_lock.release()


def last_status():
    try:
        return json.loads(db.get_setting("inbox_status") or "{}")
    except ValueError:
        return {}


def start_scheduler():
    def loop():
        time.sleep(60)
        while True:
            try:
                sync_inbox()
            except Exception as e:   # never let the loop die
                db.set_setting("inbox_status", json.dumps({"at": db.now(), "errors": [str(e)], "error_count": 1}))
            time.sleep(check_minutes() * 60)
    threading.Thread(target=loop, name="inbox-sync", daemon=True).start()


def refresh_locations():
    """List the Business Profile locations the signed-in person manages; suggest a practice by name."""
    client = cx.GBPClient(db.get_setting("gbp_refresh_token"))
    locs = client.locations()
    plist = practices()
    with db.conn() as c:
        for l in locs:
            have = c.execute("SELECT practice_id FROM gbp_locations WHERE name=?", (l["name"],)).fetchone()
            if have:
                c.execute("UPDATE gbp_locations SET title=?, address=?, account=? WHERE name=?",
                          (l["title"], l["address"], l["account"], l["name"]))
            else:
                c.execute("INSERT INTO gbp_locations(name, title, address, account, practice_id) VALUES (?,?,?,?,?)",
                          (l["name"], l["title"], l["address"], l["account"], suggest_practice(l, plist)))
    return len(locs)


def suggest_practice(loc, plist):
    words = set(re.findall(r"[a-z0-9]+", (loc["title"] + " " + loc["address"]).lower())) - {
        "orthodontics", "orthodontic", "orthodontist", "dental", "dentistry", "the", "and", "of", "smiles", "pc", "llc"}
    best, score = None, 0
    for p in plist:
        pw = set(re.findall(r"[a-z0-9]+", (p.get("campaign_filter") or p["name"]).lower())) - {"orthodontics", "the", "and"}
        s = len(words & pw)
        if s > score:
            best, score = p["id"], s
    return best


def link_locations(pairs):
    with db.conn() as c:
        for name, pid in pairs.items():
            c.execute("UPDATE gbp_locations SET practice_id=? WHERE name=?", (pid, name))


# ------------------------------------------------------------------ actions
def _platform(it, action, message):
    if it["source"] == "google":
        if action != "reply":
            raise ReviewsError("Google reviews can only be replied to.")
        cx.GBPClient(db.get_setting("gbp_refresh_token")).reply(it["external_id"], message)
        with db.conn() as c:
            c.execute("UPDATE gbp_reviews SET reply=?, reply_at=? WHERE id=?", (message, _iso(utcnow()), it["external_id"]))
        return
    tokens = _meta_tokens()[it["source"]].get(it["account_id"]) or []
    if not tokens:
        tokens = _meta_tokens(force=True)[it["source"]].get(it["account_id"]) or []
    if not tokens:
        raise ReviewsError("The hub has no access to this account's comments any more. Check the Page in Social "
                           "or the system user's assigned assets.")
    try:
        cx.comment_action(it["source"], it["external_id"], action, tokens, message)
    except cx.ConnectorError:
        # Access may have just been granted in Meta (e.g. Community activity on a Page): fetch the Page
        # tokens again, and if there's a new way in, try once more before giving up.
        fresh = _meta_tokens(force=True)[it["source"]].get(it["account_id"]) or []
        if not fresh or fresh == tokens:
            raise
        cx.comment_action(it["source"], it["external_id"], action, fresh, message)


def act(iid, action, user_name, message=""):
    """reply | like | hide | done | reopen. Returns the updated item. Platform errors raise."""
    it = item(iid)
    if not it:
        raise ReviewsError("That item isn't in the inbox any more.")
    message = (message or "").strip()
    if action == "reopen":
        with db.conn() as c:
            c.execute("UPDATE inbox SET status='open', outcome='', handled_by='', handled_at=NULL WHERE id=?", (iid,))
        return item(iid)
    if it["status"] == "done" and action != "done":
        raise ReviewsError("This one is already handled. Move it back to Needs reply first if it needs more.")
    if action == "reply":
        if not message:
            raise ReviewsError("Write a reply first.")
        if len(message) > 4000:
            raise ReviewsError("That reply is too long.")
        _platform(it, "reply", message)
    elif action in ("like", "hide"):
        if action == "like" and it["source"] != "facebook":
            raise ReviewsError("Only Facebook comments can be liked from the hub.")
        if it["source"] == "google":
            raise ReviewsError("Google reviews can't be hidden.")
        _platform(it, action, "")
    elif action != "done":
        raise ReviewsError("Unknown action.")
    outcome = {"reply": "replied", "like": "liked", "hide": "hidden", "done": "done"}[action]
    with db.conn() as c:
        c.execute("UPDATE inbox SET status='done', outcome=?, handled_by=?, handled_at=?, reply=? WHERE id=?",
                  (outcome, user_name, _iso(utcnow()), message if action == "reply" else it["reply"], iid))
    return item(iid)


def save_draft(iid, text):
    with db.conn() as c:
        c.execute("UPDATE inbox SET draft=? WHERE id=?", (text, iid))


DRAFT_RULES = (
    "You draft public replies for a group of orthodontic practices to Google reviews and Facebook / Instagram "
    "comments. The reply posts as the practice, so write as the team ('we'). Rules, strictly:\n"
    "- Never confirm or imply that the person is (or was) a patient, and never mention treatment, appointments, "
    "visits, records or anything about their care, even if they did. (US privacy law: HIPAA.)\n"
    "- Positive: thank them warmly in 1-2 short sentences; using their first name is fine.\n"
    "- Complaint: brief, sincere, not defensive; don't argue or give details; invite them to contact the office "
    "directly so the team can help. Don't promise refunds or outcomes.\n"
    "- Question (hours, insurance, prices, availability): never invent facts; invite them to call or message the "
    "office, and say the team will be happy to help.\n"
    "- Under 60 words. No hashtags, no emoji unless the comment used them, no sign-off with a staff name.\n"
    "Reply with the text of the reply only.")


def draft_reply(iid):
    from . import ai
    it = item(iid)
    if not it:
        raise ReviewsError("That item isn't in the inbox any more.")
    names = account_names()
    ctx = [f"Platform: {SOURCE_LABEL[it['source']]}", f"Business: {names.get(it['account_id'], 'the practice')}",
           f"From: {it['author'] or 'someone'}"]
    if it["stars"]:
        ctx.append(f"Star rating: {it['stars']} of 5")
    if it["post_text"]:
        ctx.append(f"Comment is on our post: {it['post_text'][:300]}")
    ctx.append(f"Their words: {it['text'][:1500] or '(no text, rating only)'}")
    text = ai._chat([{"role": "system", "content": DRAFT_RULES}, {"role": "user", "content": "\n".join(ctx)}],
                    temperature=0.6, max_tokens=400).strip().strip('"')
    save_draft(iid, text)
    return text


# ------------------------------------------------------------------ goals and monthly counts
def _goal_rows():
    with db.conn() as c:
        rows = {}
        for r in c.execute("SELECT practice_id, month, goal FROM review_goals ORDER BY month"):
            rows.setdefault(r["practice_id"], []).append((r["month"], r["goal"]))
        return rows


def goal_for(pid, month, rows=None):
    """The goal in force that month: the practice's latest goal set on or before it, else the
    'every practice' goal in force then, else the default."""
    rows = rows if rows is not None else _goal_rows()
    for key in (pid, 0):
        hit = [g for m, g in rows.get(key, []) if m <= month]
        if hit:
            return hit[-1]
    return DEFAULT_GOAL


def set_goals(goals, month=None):
    """goals: {practice id (0 = every practice): goal}. Applies from `month` (default this month) on."""
    month = month or month_key(date.today())
    with db.conn() as c:
        for pid, g in goals.items():
            if g is None:
                continue
            if g < 0 or g > 1000:
                raise ReviewsError("Goals should be between 0 and 1000.")
            c.execute("INSERT INTO review_goals(practice_id, month, goal) VALUES (?,?,?) "
                      "ON CONFLICT(practice_id, month) DO UPDATE SET goal=excluded.goal", (pid, month, int(g)))


def set_manual(month, counts):
    """counts: {practice id: count or None (clears)}."""
    if not re.fullmatch(r"\d{4}-\d{2}", month or ""):
        raise ReviewsError("Pick a month.")
    with db.conn() as c:
        for pid, n in counts.items():
            if n is None:
                c.execute("DELETE FROM review_manual WHERE practice_id=? AND month=?", (pid, month))
            elif n < 0 or n > 10000:
                raise ReviewsError("That count doesn't look right.")
            else:
                c.execute("INSERT INTO review_manual(practice_id, month, count) VALUES (?,?,?) "
                          "ON CONFLICT(practice_id, month) DO UPDATE SET count=excluded.count", (pid, month, n))


def _practice_data(plist, months):
    """Per practice: monthly counts (None = no data), rating, total, last review."""
    locs = [l for l in locations() if l["practice_id"]]
    by_p = {}
    for l in locs:
        by_p.setdefault(l["practice_id"], []).append(l)
    with db.conn() as c:
        counted = {}
        for r in c.execute("SELECT location, substr(created, 1, 7) AS m, COUNT(*) AS n FROM gbp_reviews "
                           "WHERE created >= ? GROUP BY location, m", (months[0] + "-01",)):
            counted[(r["location"], r["m"])] = r["n"]
        last = {r["location"]: r["last"] for r in c.execute("SELECT location, MAX(created) AS last FROM gbp_reviews GROUP BY location")}
        manual = {(r["practice_id"], r["month"]): r["count"] for r in c.execute("SELECT * FROM review_manual")}
    out = {}
    for p in plist:
        mine = by_p.get(p["id"], [])
        synced = [l for l in mine if l["full_at"]]
        series = []
        for m in months:
            if synced:
                series.append(sum(counted.get((l["name"], m), 0) for l in synced))
            else:
                series.append(manual.get((p["id"], m)))
        totals = [(l["rating"], l["total"]) for l in mine if l["total"]]
        n = sum(t for _, t in totals)
        lasts = [last[l["name"]] for l in synced if last.get(l["name"])]
        out[p["id"]] = {"series": series, "source": "google" if synced else ("manual" if any(v is not None for v in series) else ""),
                        "total": n or None, "rating": (sum(r * t for r, t in totals) / n) if n else None,
                        "last": max(lasts) if lasts else None, "linked": bool(mine)}
    return out


def scoreboard(practice_id=None, today=None, plist=None):
    today = today or date.today()
    plist = plist if plist is not None else practices()
    months = months_back(12, today)
    cur = months[-1]
    data = _practice_data(plist, months)
    grows = _goal_rows()
    ids = [practice_id] if practice_id else [p["id"] for p in plist]
    series, goals, reporting = [], [], []
    for i, m in enumerate(months):
        have = [pid for pid in ids if data[pid]["series"][i] is not None]
        reporting.append(len(have))
        series.append(sum(data[pid]["series"][i] for pid in have) if have else None)
        goals.append(sum(goal_for(pid, m, grows) for pid in have) if have else sum(goal_for(pid, m, grows) for pid in ids))
    top = max([v for v in series if v is not None] + goals + [1]) * 1.1
    bars = []
    for i, m in enumerate(months):
        v = series[i]
        bars.append({"label": month_label(m, short=True), "month": m, "val": v, "goal": goals[i],
                     "met": v is not None and v >= goals[i], "current": m == cur,
                     "h": 0 if v is None else round(v / top * 100, 1), "gh": round(goals[i] / top * 100, 1)})
    history = []
    for i in range(len(months) - 1, max(-1, len(months) - 7), -1):
        v, g = series[i], goals[i]
        if v is None:
            res, kind = "No count", "none"
        elif v >= g:
            res, kind = "Met" + (f" (+{v - g})" if v > g else ""), "met"
        elif months[i] == cur:
            res, kind = f"{g - v} to go", "progress"
        else:
            res, kind = f"Missed by {g - v}", "missed"
        history.append({"month": month_label(months[i]) + (" (so far)" if months[i] == cur else ""), "got": v,
                        "goal": g, "result": res, "kind": kind,
                        "partial": bool(reporting[i]) and reporting[i] < len(ids)})
    full = [(v, g) for v, g in zip(series[:-1], goals[:-1]) if v is not None]
    met_count = sum(1 for v, g in full if v >= g)
    rows = []
    now = utcnow()
    open_by_p = {}
    owners = account_practice(db.practices())
    rep_by_p, rem_by_p = {}, {}
    with db.conn() as c:
        for r in c.execute("SELECT source, account_id FROM inbox WHERE status='open'"):
            pid = owners.get((r["source"], r["account_id"]))
            open_by_p[pid] = open_by_p.get(pid, 0) + 1
        for r in c.execute("SELECT account_id, status, outcome FROM inbox WHERE source='google' AND reported_at IS NOT NULL"):
            pid = owners.get(("google", r["account_id"]))
            if r["status"] == "reported":
                rep_by_p[pid] = rep_by_p.get(pid, 0) + 1
            elif r["outcome"] == "removed":
                rem_by_p[pid] = rem_by_p.get(pid, 0) + 1
    for p in plist:
        d = data[p["id"]]
        got = d["series"][-1]
        g = goal_for(p["id"], cur, grows)
        last6 = [(d["series"][i], goal_for(p["id"], months[i], grows)) for i in range(5, 11)]
        days = (now - _dt(d["last"])).days if d["last"] else None
        rows.append({"id": p["id"], "name": p["name"], "rating": d["rating"], "total": d["total"], "got": got, "goal": g,
                     "pct": min(100, round((got or 0) / g * 100)) if g else 100, "met": got is not None and got >= g,
                     "met6": sum(1 for v, gg in last6 if v is not None and v >= gg),
                     "known6": sum(1 for v, _ in last6 if v is not None),
                     "last": ago(d["last"], now) if d["last"] else None, "days": days,
                     "waiting": open_by_p.get(p["id"], 0), "source": d["source"], "linked": d["linked"],
                     "reported": rep_by_p.get(p["id"], 0), "removed": rem_by_p.get(p["id"], 0),
                     "avg6": (sum(v for v, _ in last6 if v is not None) / max(1, sum(1 for v, _ in last6 if v is not None)))
                     if any(v is not None for v, _ in last6) else None})
    sel = [r for r in rows if not practice_id or r["id"] == practice_id]
    rated = [(r["rating"], r["total"]) for r in sel if r["rating"] and r["total"]]
    n = sum(t for _, t in rated)
    return {"months": months, "bars": bars, "history": history, "rows": rows, "series": series, "goals": goals,
            "this": series[-1], "this_goal": goals[-1], "last": series[-2], "last_goal": goals[-2],
            "met_count": met_count, "full_months": len(full),
            "at_goal": sum(1 for r in sel if r["met"]), "reporting": reporting[-1], "of": len(ids),
            "rating": (sum(r * t for r, t in rated) / n) if n else None, "total": n or None,
            "quiet": sum(1 for r in sel if r["days"] is not None and r["days"] > 30),
            "reported": sum(r["reported"] for r in sel), "removed": sum(r["removed"] for r in sel),
            "current_label": month_label(cur), "default_goal": goal_for(0, cur, grows)}


def export_xlsx(plist=None, today=None):
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    plist = plist if plist is not None else practices()
    months = months_back(12, today)
    data = _practice_data(plist, months)
    grows = _goal_rows()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Google reviews"
    ws.append([f"New Google reviews by month · {month_label(months[0])} to {month_label(months[-1])}"])
    ws["A1"].font = Font(bold=True, size=13)
    head = ["Practice"]
    for m in months:
        head += [f"{month_label(m)} new", "goal"]
    head += ["Rating", "Total reviews"]
    ws.append(head)
    for c in ws[2]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0F7471")
    for p in plist:
        d = data[p["id"]]
        row = [p["name"]]
        for i, m in enumerate(months):
            row += [d["series"][i], goal_for(p["id"], m, grows)]
        row += [round(d["rating"], 2) if d["rating"] else None, d["total"]]
        ws.append(row)
    ws.column_dimensions["A"].width = 34
    ws.freeze_panes = "B3"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
