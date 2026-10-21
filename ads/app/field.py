"""Field visits: marketing agents log visits to dental offices, with spend against budgets and routes.

Google's terms: only place IDs are kept from Google. An office's name and address in our
records come from the NPI Registry (public data) or from what the agent types; Google's own
name/address are only ever shown live on the map. Visit details, spend, receipts and the
phone's check-in location are our own data and are kept.
"""
from __future__ import annotations

import base64
import calendar
import io
import json
import math
import os
import re
import secrets
import urllib.parse
from datetime import date, datetime, timedelta

from . import db, places, referrals, users

RATINGS = ("cold", "warm", "hot", "referring")
RATING_LABELS = {"cold": "Cold", "warm": "Warm", "hot": "Hot", "referring": "Referring"}
MET_ROLES = ("Office manager", "Doctor", "Front desk", "Hygienist / assistant", "Other")
DROP_CHIPS = ("Lunch", "Treats", "Referral pads", "Swag", "Brochures")
MAX_PEOPLE = 12
EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[A-Za-z]{2,}$")
ON_SITE_MILES = 0.2
MAX_RECEIPT = 8 * 1024 * 1024
MAPS_STOPS_PER_LINK = 10   # Google Maps app: origin + up to 9 waypoints + destination

SCHEMA = """
CREATE TABLE IF NOT EXISTS field_offices (
  id INTEGER PRIMARY KEY,
  place_id TEXT UNIQUE,
  npi_key TEXT,
  name TEXT NOT NULL DEFAULT '',
  address TEXT DEFAULT '',
  lat REAL, lng REAL,
  our_office_id INTEGER,
  created_at TEXT NOT NULL,
  created_by TEXT DEFAULT ''
);
CREATE INDEX IF NOT EXISTS field_offices_npi ON field_offices(npi_key);
CREATE TABLE IF NOT EXISTS visits (
  id INTEGER PRIMARY KEY,
  office_id INTEGER NOT NULL REFERENCES field_offices(id) ON DELETE CASCADE,
  user_id INTEGER,
  user_name TEXT DEFAULT '',
  visited_at TEXT NOT NULL,
  met_name TEXT DEFAULT '',
  met_role TEXT DEFAULT '',
  rating TEXT DEFAULT '',
  dropped TEXT DEFAULT '',
  cost REAL DEFAULT 0,
  receipt TEXT DEFAULT '',
  followup TEXT DEFAULT '',
  notes TEXT DEFAULT '',
  phone_lat REAL, phone_lng REAL, phone_acc REAL,
  distance_mi REAL,
  on_site INTEGER,
  created_at TEXT NOT NULL
);
CREATE INDEX IF NOT EXISTS visits_when ON visits(visited_at);
CREATE INDEX IF NOT EXISTS visits_office ON visits(office_id, visited_at);
CREATE TABLE IF NOT EXISTS visit_people (
  id INTEGER PRIMARY KEY,
  visit_id INTEGER NOT NULL REFERENCES visits(id) ON DELETE CASCADE,
  pos INTEGER NOT NULL DEFAULT 0,
  name TEXT DEFAULT '',
  email TEXT DEFAULT '',
  role TEXT DEFAULT ''
);
CREATE INDEX IF NOT EXISTS visit_people_visit ON visit_people(visit_id);
CREATE TABLE IF NOT EXISTS field_links (
  map_office_id INTEGER PRIMARY KEY,
  practice_id INTEGER
);
"""


class FieldError(ValueError):
    pass


def init():
    with db.conn() as c:
        c.executescript(SCHEMA)
    os.makedirs(receipt_dir(), exist_ok=True)


def receipt_dir():
    return os.path.join(os.path.dirname(db.DB_PATH) or ".", "receipts")


# ------------------------------------------------------------------ helpers
def _f(v):
    try:
        v = float(v)
        return v if math.isfinite(v) else None
    except (TypeError, ValueError):
        return None


def money_in(text):
    """'$12.50' / '12' / '' -> float (0 when blank)."""
    t = re.sub(r"[^\d.\-]", "", str(text or ""))
    if not t:
        return 0.0
    v = _f(t)
    if v is None or v < 0 or v > 100000:
        raise FieldError("That amount doesn't look right.")
    return round(v, 2)


def _clean(s, n=500):
    return re.sub(r"\s+", " ", str(s or "")).strip()[:n]


def _date_or_blank(s):
    s = (s or "").strip()[:10]
    if not s:
        return ""
    try:
        return date.fromisoformat(s).isoformat()
    except ValueError:
        raise FieldError("That follow-up date isn't valid.")


def _when(s):
    s = (s or "").strip()
    if not s:
        return datetime.now().isoformat(timespec="minutes")
    try:
        return datetime.fromisoformat(s[:16]).isoformat(timespec="minutes")
    except ValueError:
        raise FieldError("That visit date and time isn't valid.")


def nearest_office(lat, lng):
    best = None
    for o in places.offices():
        if o["lat"] is None:
            continue
        d = places.miles_between(o["lat"], o["lng"], lat, lng)
        if best is None or d < best[1]:
            best = (o["id"], d)
    return best[0] if best else None


# ------------------------------------------------------------------ offices
def _npi(key):
    with db.conn() as c:
        r = c.execute("SELECT * FROM ref_practices WHERE key=?", (key,)).fetchone()
    return dict(r) if r else None


def _npi_for_place(place_id):
    with db.conn() as c:
        r = c.execute("SELECT * FROM ref_practices WHERE place_id=? AND place_id!='' LIMIT 1", (place_id,)).fetchone()
    return dict(r) if r else None


def _map_place(place_id):
    with db.conn() as c:
        r = c.execute("SELECT * FROM map_places WHERE place_id=?", (place_id,)).fetchone()
    return dict(r) if r else None


def _npi_address(r):
    return ", ".join(x for x in (r.get("address"), r.get("city"), f"{r.get('state') or ''} {r.get('zip') or ''}".strip()) if x)


def office(fid):
    with db.conn() as c:
        r = c.execute("SELECT * FROM field_offices WHERE id=?", (fid,)).fetchone()
    return _with_coords(dict(r)) if r else None


def _with_coords(o):
    """Coordinates for distance checks and routes: Google pins use the map's current (30-day) location."""
    o["here_lat"], o["here_lng"] = o.get("lat"), o.get("lng")
    if o.get("place_id"):
        p = _map_place(o["place_id"])
        if p and p["lat"] is not None:
            o["here_lat"], o["here_lng"] = p["lat"], p["lng"]
            if o.get("our_office_id") is None:
                o["our_office_id"] = p["office_id"]
    elif o.get("npi_key") and o["here_lat"] is None:
        n = _npi(o["npi_key"])
        if n:
            o["here_lat"], o["here_lng"] = n["lat"], n["lng"]
    return o


def find_office(place_id=None, npi_key=None):
    """The field record for a map pin or NPI practice, if one exists yet (not created)."""
    with db.conn() as c:
        if place_id:
            r = c.execute("SELECT * FROM field_offices WHERE place_id=?", (place_id,)).fetchone()
            if r:
                return _with_coords(dict(r))
        if npi_key:
            r = c.execute("SELECT * FROM field_offices WHERE npi_key=? ORDER BY id LIMIT 1", (npi_key,)).fetchone()
            if r:
                return _with_coords(dict(r))
    return None


def office_draft(place_id=None, npi_key=None):
    """What the log form shows for an office not visited before: name/address from NPI when known."""
    existing = find_office(place_id, npi_key)
    if existing:
        return existing
    if npi_key:
        n = _npi(npi_key)
        if not n:
            raise FieldError("That practice isn't in the list any more.")
        if n.get("place_id"):
            ex = find_office(n["place_id"])
            if ex:
                return ex
        return _with_coords({"id": None, "place_id": n.get("place_id") or None, "npi_key": n["key"], "name": n["name"],
                             "address": _npi_address(n), "lat": n["lat"], "lng": n["lng"], "our_office_id": None})
    if place_id:
        if not places._valid_place_id(place_id):
            raise FieldError("That map pin isn't valid.")
        n = _npi_for_place(place_id)
        return _with_coords({"id": None, "place_id": place_id, "npi_key": n["key"] if n else None,
                             "name": n["name"] if n else "", "address": _npi_address(n) if n else "",
                             "lat": None, "lng": None, "our_office_id": None})
    return {"id": None, "place_id": None, "npi_key": None, "name": "", "address": "", "lat": None, "lng": None,
            "here_lat": None, "here_lng": None, "our_office_id": None}


def ensure_office(place_id=None, npi_key=None, name="", address="", lat=None, lng=None, by="", office_id=None):
    """Find or create the field record; the name is ours (NPI or typed), never Google's."""
    name, address = _clean(name, 160), _clean(address, 300)
    o = office(int(office_id)) if str(office_id or "").isdigit() else None
    o = o or find_office(place_id, npi_key)
    if o:
        if name and name != o["name"]:
            with db.conn() as c:
                c.execute("UPDATE field_offices SET name=? WHERE id=?", (name, o["id"]))
            o["name"] = name
        return o
    if not name:
        raise FieldError("Add the office's name (as on their sign) so it reads well in reports.")
    lat, lng = _f(lat), _f(lng)
    if place_id:
        if not places._valid_place_id(place_id):
            raise FieldError("That map pin isn't valid.")
        lat = lng = None   # Google coordinates aren't ours to keep; the map's own copy is used
        p = _map_place(place_id)
        our = p["office_id"] if p else None
    elif npi_key:
        n = _npi(npi_key)
        if not n:
            raise FieldError("That practice isn't in the list any more.")
        lat, lng = n["lat"], n["lng"]
        our = nearest_office(lat, lng) if lat is not None else None
    else:
        if lat is None or lng is None:
            if address:
                got = geocode_one(address)
                if got:
                    lat, lng = got
        our = nearest_office(lat, lng) if lat is not None else None
    with db.conn() as c:
        cur = c.execute("INSERT INTO field_offices(place_id, npi_key, name, address, lat, lng, our_office_id, created_at, "
                        "created_by) VALUES (?,?,?,?,?,?,?,?,?)",
                        (place_id or None, npi_key or None, name, address, lat, lng, our, db.now(), by))
        fid = cur.lastrowid
    return office(fid)


def geocode_one(address):
    """Free-text US address -> (lat, lng) with the Census geocoder (free, not Google), or None."""
    url = referrals.GEOCODER_URL.replace("addressbatch", "onelineaddress") + "?" + urllib.parse.urlencode(
        {"address": address, "benchmark": "Public_AR_Current", "format": "json"})
    try:
        data = json.loads(referrals._raw(url, timeout=20))
        m = data["result"]["addressMatches"]
        if m:
            return float(m[0]["coordinates"]["y"]), float(m[0]["coordinates"]["x"])
    except (referrals.RefError, ValueError, KeyError, TypeError, IndexError):
        pass
    return None


# ------------------------------------------------------------------ people met
def people_in(form):
    """The people met on a visit, from form["people"] (list of dicts) or the older single met_name/met_role pair.
    Blank rows are dropped; a bad email raises FieldError."""
    raw = form.get("people")
    if raw is None:
        raw = [{"name": form.get("met_name"), "email": form.get("met_email"), "role": form.get("met_role")}]
    out = []
    for p in raw:
        name, email, role = _clean(p.get("name"), 120), _clean(p.get("email"), 160).lower(), _clean(p.get("role"), 60)
        if not name and not email:
            continue
        if email and not EMAIL_RE.match(email):
            raise FieldError(f"“{email}” doesn't look like an email address.")
        out.append({"name": name, "email": email, "role": role})
    if len(out) > MAX_PEOPLE:
        raise FieldError(f"That's more than {MAX_PEOPLE} people; log the main contacts.")
    return out


def with_people(rows):
    """Adds r["people"] to each visit row. Visits logged before people had their own table fall back to met_name."""
    ids = [r["id"] for r in rows]
    found = {}
    if ids:
        with db.conn() as c:
            for i in range(0, len(ids), 500):
                chunk = ids[i:i + 500]
                for p in c.execute(f"SELECT visit_id, name, email, role FROM visit_people WHERE visit_id IN "
                                   f"({','.join('?' * len(chunk))}) ORDER BY visit_id, pos", chunk):
                    found.setdefault(p["visit_id"], []).append({"name": p["name"], "email": p["email"], "role": p["role"]})
    for r in rows:
        r["people"] = found.get(r["id"]) or ([{"name": r["met_name"], "email": "", "role": r.get("met_role") or ""}]
                                             if r.get("met_name") else [])
    return rows


def person_label(p):
    return p["name"] or p["email"]


# ------------------------------------------------------------------ visits
def log_visit(user, form, receipt_bytes=None, receipt_type=""):
    """form: dict of the log form's fields. Returns the visit id."""
    o = ensure_office(form.get("place_id") or None, form.get("npi_key") or None, form.get("name"),
                      form.get("address"), form.get("office_lat"), form.get("office_lng"), user.get("name", ""),
                      form.get("office_id"))
    rating = (form.get("rating") or "").strip().lower()
    if rating and rating not in RATINGS:
        raise FieldError("Pick how warm the relationship is.")
    cost = money_in(form.get("cost"))
    plat, plng, pacc = _f(form.get("phone_lat")), _f(form.get("phone_lng")), _f(form.get("phone_acc"))
    dist = on_site = None
    if plat is not None and plng is not None and o.get("here_lat") is not None:
        dist = round(places.miles_between(plat, plng, o["here_lat"], o["here_lng"]), 2)
        slack = min((pacc or 0) / 1609.34, 0.3)   # allow for the phone's stated accuracy (capped)
        on_site = 1 if dist <= ON_SITE_MILES + slack else 0
    elif plat is not None and o.get("here_lat") is None and not o.get("place_id") and not o.get("npi_key"):
        # A brand-new office logged "where I am": the phone's location is the office.
        with db.conn() as c:
            c.execute("UPDATE field_offices SET lat=?, lng=?, our_office_id=COALESCE(our_office_id, ?) WHERE id=?",
                      (plat, plng, nearest_office(plat, plng), o["id"]))
        dist, on_site = 0.0, 1
    people = people_in(form)
    met_name = _clean(", ".join(person_label(p) for p in people), 300)
    met_role = _clean(", ".join(dict.fromkeys(p["role"] for p in people if p["role"])), 200)
    receipt = save_receipt(receipt_bytes, receipt_type) if receipt_bytes else ""
    with db.conn() as c:
        cur = c.execute(
            "INSERT INTO visits(office_id, user_id, user_name, visited_at, met_name, met_role, rating, dropped, cost, "
            "receipt, followup, notes, phone_lat, phone_lng, phone_acc, distance_mi, on_site, created_at) "
            "VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
            (o["id"], user.get("id") or None, user.get("name", ""), _when(form.get("visited_at")),
             met_name, met_role, rating,
             _clean(form.get("dropped"), 300), cost, receipt, _date_or_blank(form.get("followup")),
             str(form.get("notes") or "").strip()[:4000], plat, plng, pacc, dist, on_site, db.now()))
        vid = cur.lastrowid
        c.executemany("INSERT INTO visit_people(visit_id, pos, name, email, role) VALUES (?,?,?,?,?)",
                      [(vid, i, p["name"], p["email"], p["role"]) for i, p in enumerate(people)])
        return vid


def visit(vid):
    with db.conn() as c:
        r = c.execute("SELECT v.*, o.name AS office_name, o.place_id, o.npi_key, o.address AS office_address "
                      "FROM visits v JOIN field_offices o ON o.id=v.office_id WHERE v.id=?", (vid,)).fetchone()
    return with_people([dict(r)])[0] if r else None


def delete_visit(vid):
    v = visit(vid)
    if not v:
        return
    with db.conn() as c:
        c.execute("DELETE FROM visit_people WHERE visit_id=?", (vid,))
        c.execute("DELETE FROM visits WHERE id=?", (vid,))
    if v["receipt"]:
        try:
            os.remove(os.path.join(receipt_dir(), v["receipt"]))
        except OSError:
            pass


def save_receipt(data, content_type=""):
    if len(data) > MAX_RECEIPT:
        raise FieldError("That photo is too large (8 MB at most).")
    kind = _image_kind(data)
    if not kind:
        raise FieldError("The receipt needs to be a photo (JPEG, PNG, WebP or HEIC).")
    name = f"{datetime.now():%Y%m}-{secrets.token_hex(8)}.{kind}"
    with open(os.path.join(receipt_dir(), name), "wb") as f:
        f.write(data)
    return name


def _image_kind(data):
    if data[:3] == b"\xff\xd8\xff":
        return "jpg"
    if data[:8] == b"\x89PNG\r\n\x1a\n":
        return "png"
    if data[:4] == b"RIFF" and data[8:12] == b"WEBP":
        return "webp"
    if data[4:12] in (b"ftypheic", b"ftypheix", b"ftypmif1", b"ftyphevc"):
        return "heic"
    return None


RECEIPT_NAME = re.compile(r"^\d{6}-[0-9a-f]{16}\.(jpg|png|webp|heic)$")


def receipt_path(name):
    if not RECEIPT_NAME.match(name or ""):
        return None
    p = os.path.join(receipt_dir(), name)
    return p if os.path.exists(p) else None


def read_receipt(data):
    """Ask the AI for the receipt's total. Returns {"total": float|None, "merchant": str}."""
    from . import ai
    kind = _image_kind(data)
    if kind not in ("jpg", "png", "webp"):
        raise FieldError("The AI can read JPEG, PNG or WebP photos; type the amount for this one.")
    mime = {"jpg": "image/jpeg", "png": "image/png", "webp": "image/webp"}[kind]
    url = f"data:{mime};base64,{base64.b64encode(data).decode()}"
    text = ai._chat([
        {"role": "system", "content": "You read store and restaurant receipts. Reply with one JSON object only."},
        {"role": "user", "content": [
            {"type": "text", "text": 'Read this receipt. Reply as {"total": <the final amount paid including tax and tip, '
                                     'as a number, or null if you cannot read it>, "merchant": "<store name or empty>"}'},
            {"type": "image_url", "image_url": {"url": url}}]},
    ], temperature=0, max_tokens=200, json_mode=True)
    try:
        raw = text[text.find("{"):text.rfind("}") + 1]
        got = json.loads(raw)
    except ValueError:
        raise FieldError("The AI couldn't read that receipt; type the amount instead.")
    total = _f(got.get("total"))
    return {"total": round(total, 2) if total is not None and 0 <= total < 100000 else None,
            "merchant": _clean(got.get("merchant"), 80)}


# ------------------------------------------------------------------ periods and activity
def period_bounds(kind, day):
    """('day'|'week'|'month', date) -> (start, end_exclusive, label, prev_day, next_day)."""
    if kind == "day":
        start, end = day, day + timedelta(days=1)
        label = day.strftime("%A, %B %-d, %Y")
        prev_, next_ = day - timedelta(days=1), day + timedelta(days=1)
    elif kind == "week":
        start = day - timedelta(days=day.weekday())   # Monday
        end = start + timedelta(days=7)
        last = end - timedelta(days=1)
        label = f"{start:%b %-d} – {last:%b %-d, %Y}" if start.month != last.month else f"{start:%b %-d}–{last:%-d, %Y}"
        prev_, next_ = start - timedelta(days=7), start + timedelta(days=7)
    else:
        kind = "month"
        start = day.replace(day=1)
        end = (start + timedelta(days=32)).replace(day=1)
        label = start.strftime("%B %Y")
        prev_, next_ = (start - timedelta(days=1)).replace(day=1), end
    return start, end, label, prev_, next_


def visits_between(start, end, user_id=None, practice_id=None):
    with db.conn() as c:
        rows = [dict(r) for r in c.execute(
            "SELECT v.*, o.name AS office_name, o.place_id, o.npi_key, o.address AS office_address, "
            "o.our_office_id, o.lat AS office_lat, o.lng AS office_lng FROM visits v "
            "JOIN field_offices o ON o.id=v.office_id WHERE v.visited_at >= ? AND v.visited_at < ? "
            "ORDER BY v.visited_at DESC, v.id DESC", (start.isoformat(), end.isoformat()))]
    links = practice_links()
    for r in rows:
        if r["our_office_id"] is None and r["place_id"]:
            p = _map_place(r["place_id"])
            r["our_office_id"] = p["office_id"] if p else None
        r["practice_id"] = links.get(r["our_office_id"])
    if user_id:
        rows = [r for r in rows if r["user_id"] == user_id]
    if practice_id:
        rows = [r for r in rows if r["practice_id"] == practice_id]
    return with_people(rows)


def pacing(spend, budget, start, today=None):
    """Same shape as the ads pacing ruler, for one calendar month of field spend."""
    today = today or date.today()
    dim = calendar.monthrange(start.year, start.month)[1]
    if (today.year, today.month) == (start.year, start.month):
        elapsed = today.day
    elif today < start:
        elapsed = 0
    else:
        elapsed = dim
    res = {"available": True, "month": start.strftime("%B %Y"), "days_elapsed": elapsed, "days_in_month": dim,
           "spend": round(spend, 2), "budget": budget or None,
           "projected": round(spend / elapsed * dim, 2) if elapsed else spend}
    if budget:
        expected = round(budget * elapsed / dim, 2)
        ratio = spend / expected if expected else None
        res.update(expected=expected, ratio=ratio, remaining=round(budget - spend, 2),
                   daily_to_hit=round((budget - spend) / (dim - elapsed), 2) if dim > elapsed else 0.0,
                   status="unknown" if ratio is None else "over" if ratio > 1.10 else "under" if ratio < 0.90 else "on-track")
    else:
        res["status"] = "no-budget"
    return res


def summary(rows):
    return {"visits": len(rows), "offices": len({r["office_id"] for r in rows}),
            "spend": round(sum(r["cost"] or 0 for r in rows), 2),
            "on_site": sum(1 for r in rows if r["on_site"] == 1),
            "checked": sum(1 for r in rows if r["on_site"] is not None),
            "agents": len({r["user_name"] for r in rows})}


def by_day(rows):
    """[(date, [visits oldest first])] newest day first, for the 'where they've been' timeline."""
    days = {}
    for r in rows:
        days.setdefault(r["visited_at"][:10], []).append(r)
    return [(date.fromisoformat(d), sorted(v, key=lambda r: r["visited_at"])) for d, v in sorted(days.items(), reverse=True)]


def budgets(month_start, rows_month, today=None):
    """Pacing per practice (spend on offices linked to it) and per agent, for one month."""
    pr = {p["id"]: p for p in db.practices()}
    by_practice, by_agent = {}, {}
    for r in rows_month:
        if r["practice_id"]:
            by_practice[r["practice_id"]] = by_practice.get(r["practice_id"], 0) + (r["cost"] or 0)
        key = r["user_id"] or r["user_name"]
        by_agent[key] = by_agent.get(key, 0) + (r["cost"] or 0)
    unlinked = round(sum(r["cost"] or 0 for r in rows_month if not r["practice_id"]), 2)
    practices_out = []
    for pid, p in pr.items():
        b = p.get("field_budget") or 0
        if b or by_practice.get(pid):
            practices_out.append({"id": pid, "name": p["name"],
                                  "pacing": pacing(by_practice.get(pid, 0), b, month_start, today)})
    agents_out = []
    people = {u["id"]: u for u in users.all_users()}
    for uid, u in people.items():
        b = u.get("field_budget") or 0
        if b or by_agent.get(uid) or (u.get("field") and u["active"]):
            agents_out.append({"id": uid, "name": u["name"], "pacing": pacing(by_agent.get(uid, 0), b, month_start, today)})
    for key, spend in by_agent.items():
        if isinstance(key, str):   # visits by someone no longer on the team
            agents_out.append({"id": None, "name": key, "pacing": pacing(spend, 0, month_start, today)})
    practices_out.sort(key=lambda x: -x["pacing"]["spend"])
    agents_out.sort(key=lambda x: (-x["pacing"]["spend"], x["name"].lower()))
    return {"practices": practices_out, "agents": agents_out, "unlinked": unlinked}


def office_status():
    """Latest visit per office: rating, when, by whom, open follow-up. Keyed by field office id."""
    with db.conn() as c:
        rows = [dict(r) for r in c.execute(
            "SELECT v.office_id, v.visited_at, v.user_id, v.user_name, v.rating, v.followup, v.met_name, "
            "(SELECT COUNT(*) FROM visits x WHERE x.office_id=v.office_id) AS n, "
            "o.place_id, o.npi_key, o.name, o.address, o.lat, o.lng "
            "FROM visits v JOIN field_offices o ON o.id=v.office_id "
            "WHERE v.id = (SELECT id FROM visits y WHERE y.office_id=v.office_id ORDER BY visited_at DESC, id DESC LIMIT 1)")]
    return rows


def followups(today=None, within_days=7, user_id=None):
    today = today or date.today()
    soon = (today + timedelta(days=within_days)).isoformat()
    out = []
    for s in office_status():
        if s["followup"] and s["followup"] <= soon:
            o = _with_coords({"place_id": s["place_id"], "npi_key": s["npi_key"], "lat": s["lat"], "lng": s["lng"]})
            out.append({**s, "overdue": s["followup"] < today.isoformat(), "due_today": s["followup"] == today.isoformat(),
                        "here_lat": o["here_lat"], "here_lng": o["here_lng"]})
    if user_id:
        out = [f for f in out if f["user_id"] == user_id]
    return sorted(out, key=lambda f: f["followup"])


def map_overlay():
    """For the map page: visit status per Google pin, plus visited offices that aren't Google pins."""
    by_place, extra = {}, []
    today = date.today().isoformat()
    for s in office_status():
        info = {"rating": s["rating"], "last": s["visited_at"][:10], "by": s["user_name"], "n": s["n"],
                "due": s["followup"] or None, "overdue": bool(s["followup"] and s["followup"] < today)}
        if s["place_id"]:
            by_place[s["place_id"]] = info
        else:
            o = _with_coords({"place_id": None, "npi_key": s["npi_key"], "lat": s["lat"], "lng": s["lng"]})
            if o["here_lat"] is not None:
                extra.append({**info, "fid": s["office_id"], "name": s["name"], "address": s["address"],
                              "npi": s["npi_key"], "lat": o["here_lat"], "lng": o["here_lng"]})
    return {"visits": by_place, "visited_offices": extra}


# ------------------------------------------------------------------ practice links and budgets
def _words(s):
    return [w for w in re.findall(r"[a-z0-9]+", (s or "").lower()) if len(w) > 2 and w not in
            {"the", "and", "orthodontics", "orthodontic", "ortho", "dental", "dentistry", "family", "smiles", "llc", "pllc"}]


def suggest_practice(office_label, practices_list):
    """Which practice an office belongs to, from the names: 'Heights – Defiance' -> Heights Orthodontics."""
    label = set(_words(office_label))
    best = None
    for p in practices_list:
        words = _words(p["name"]) + [w for part in (p.get("campaign_filter") or "").split(",") for w in _words(part)
                                    if not part.strip().startswith("-")]
        score = sum(1 for w in set(words) if w in label)
        if score and (best is None or score > best[0]):
            best = (score, p["id"])
    return best[1] if best else None


def practice_links():
    """{map office id: practice id}: saved choices, otherwise the suggestion from the names."""
    with db.conn() as c:
        saved = {r["map_office_id"]: r["practice_id"] for r in c.execute("SELECT * FROM field_links")}
    plist = db.practices()
    out = {}
    for o in places.offices():
        out[o["id"]] = saved[o["id"]] if o["id"] in saved else suggest_practice(o["label"], plist)
    return out


def set_links(pairs):
    with db.conn() as c:
        for oid, pid in pairs.items():
            c.execute("INSERT INTO field_links(map_office_id, practice_id) VALUES (?,?) "
                      "ON CONFLICT(map_office_id) DO UPDATE SET practice_id=excluded.practice_id", (oid, pid))


def set_practice_budget(pid, amount):
    with db.conn() as c:
        c.execute("UPDATE practices SET field_budget=? WHERE id=?", (amount, pid))


def set_agent_budget(uid, amount):
    with db.conn() as c:
        c.execute("UPDATE users SET field_budget=? WHERE id=?", (amount, uid))


# ------------------------------------------------------------------ routes
def order_stops(start, stops):
    """Nearest-neighbour then 2-opt by straight-line distance; start is (lat, lng) or None.
    stops: [{"lat", "lng", ...}]. Returns the stops reordered (an open path, no return trip)."""
    pts = [s for s in stops if _f(s.get("lat")) is not None and _f(s.get("lng")) is not None]
    rest = [s for s in stops if s not in pts]
    if len(pts) < 2:
        return pts + rest

    def d(a, b):
        return places.miles_between(a[0], a[1], b[0], b[1])
    xy = [(float(s["lat"]), float(s["lng"])) for s in pts]
    left = list(range(len(pts)))
    route = []
    cur = start
    if cur is None:   # no start: begin at the stop furthest out, so the path sweeps across
        far = max(left, key=lambda i: sum(d(xy[i], xy[j]) for j in left))
        route.append(far); left.remove(far); cur = xy[far]
    while left:
        nxt = min(left, key=lambda i: d(cur, xy[i]))
        route.append(nxt); left.remove(nxt); cur = xy[nxt]

    def length(r):
        total = d(start, xy[r[0]]) if start else 0
        return total + sum(d(xy[r[i]], xy[r[i + 1]]) for i in range(len(r) - 1))
    improved = True
    while improved:
        improved = False
        for i in range(len(route) - 1):
            for k in range(i + 1, len(route)):
                cand = route[:i] + route[i:k + 1][::-1] + route[k + 1:]
                if length(cand) + 1e-9 < length(route):
                    route, improved = cand, True
    return [pts[i] for i in route] + rest


def maps_links(origin, stops):
    """Google Maps app links for driving the stops in order; split into legs of up to 10 stops.
    origin: 'lat,lng' text, an address, or '' for the phone's current location."""
    links = []
    i, here = 0, origin
    while i < len(stops):
        leg = stops[i:i + MAPS_STOPS_PER_LINK]
        dest = leg[-1]
        q = {"api": "1", "travelmode": "driving", "destination": f"{dest['lat']},{dest['lng']}"}
        if dest.get("place_id"):
            q["destination_place_id"] = dest["place_id"]
        if here:
            q["origin"] = here
        if len(leg) > 1:
            q["waypoints"] = "|".join(f"{s['lat']},{s['lng']}" for s in leg[:-1])
        links.append({"url": "https://www.google.com/maps/dir/?" + urllib.parse.urlencode(q, safe=",|"),
                      "from": i + 1, "to": i + len(leg)})
        here = f"{dest['lat']},{dest['lng']}"
        i += len(leg)
    return links


# ------------------------------------------------------------------ export
def export_xlsx(rows, title):
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Visits"
    head = ["Date", "Time", "Agent", "Office", "Address", "Met", "Role", "Email", "Relationship", "Dropped off", "Cost",
            "On site", "Miles from office", "Follow-up", "Notes", "Receipt"]
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(head)
    for c in ws[2]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0F7471")
    for r in sorted(rows, key=lambda r: r["visited_at"]):
        ppl = r.get("people") or []
        ws.append([r["visited_at"][:10], r["visited_at"][11:16], r["user_name"], r["office_name"], r["office_address"],
                   "; ".join(person_label(p) for p in ppl), "; ".join(p["role"] or "–" for p in ppl) if any(p["role"] for p in ppl) else "",
                   "; ".join(p["email"] or "–" for p in ppl) if any(p["email"] for p in ppl) else "",
                   RATING_LABELS.get(r["rating"], ""), r["dropped"], r["cost"] or 0,
                   {1: "Yes", 0: "No"}.get(r["on_site"], "Not shared"), r["distance_mi"], r["followup"], r["notes"],
                   "Yes" if r["receipt"] else ""])
        ws.cell(ws.max_row, 11).number_format = '"$"#,##0.00'
    total_row = ws.max_row + 1
    ws.cell(total_row, 10, "Total").font = Font(bold=True)
    ws.cell(total_row, 11, f"=SUM(K3:K{total_row - 1})").number_format = '"$"#,##0.00'
    for col, w in zip("ABCDEFGHIJKLMNOP", (11, 7, 16, 30, 34, 22, 18, 28, 12, 22, 10, 9, 10, 11, 40, 8)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A3"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
